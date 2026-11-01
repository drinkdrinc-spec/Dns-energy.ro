import io
from html import escape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
import docgen as D
import cef as CEF

M = D.M


def first_present(*values):
    return next((value for value in values if value not in (None, "")), None)
OP_NAMES = {"deer": "Distribuție Energie Electrică Romania S.A. (DEER)", "delgaz": "Delgaz Grid S.A.", "rer": "Rețele Electrice România S.A.",
            "oltenia": "Distribuție Energie Oltenia S.A."}
DOCS = [("plan", "Plan de situație la scară, cu viza emitentului certificatului de urbanism (sau schiță de amplasament pentru construcție existentă)", True),
        ("cf", "Extras de carte funciară actualizat (CF / nr. cadastral)", True),
        ("prop", "Act de proprietate sau alt înscris care atestă dreptul de folosință asupra terenului / incintei / clădirii", True),
        ("id", "Act de identitate / certificat de înregistrare la registrul comerțului / alte acte legale de funcționare", True),
        ("cu", "Certificat de urbanism în termen de valabilitate (construcție nouă / extindere)", False),
        ("ac", "Autorizație de construire (organizare de șantier / obiectiv)", False),
        ("puz", "PUZ / PUD aprobat, dacă a fost solicitat prin certificatul de urbanism", False),
        ("acord", "Acordul notarial al proprietarului, dacă solicitantul nu este proprietar", False),
        ("imp", "Împuternicire pentru reprezentant / proiectant (dacă cererea se depune prin împuternicit)", False),
        ("gdpr", "Notă oficială DEER privind prelucrarea datelor personale (după caz)", False),
        ("authenticity", "Declarație DEER privind autenticitatea datelor (obligatorie pentru toate tipurile de cereri)", "deer"),
        ("memoriu", "Date tehnice și energetice (memoriu explicativ)", False),
        ("schema", "Schema electrică monofilară a instalației de producere și modul de racordare (loc de consum și producere)", "prod"),
        ("fise", "Fișe tehnice / certificate de conformitate ale invertoarelor și unităților generatoare (loc de consum și producere)", "prod"),
        ("pram", "Buletin de încercare a prizei de pământ", False),
        ("receptie", "Proces-verbal de recepție a instalației de producere", False),
        ("constatator", "Certificat constatator ONRC (prosumator persoană juridică)", False),
        ("atr_prev", "Copie aviz tehnic / certificat de racordare obținut anterior (spor de putere)", False)]
TYPES = [("nou", "utilizator nou definitiv"), ("temp", "utilizator nou temporar"), ("spor", "dezvoltarea unui utilizator existent (spor de putere)"),
         ("mod_dep", "modificarea elementelor de natură tehnică cu depășirea puterii aprobate anterior"),
         ("mod_fara", "modificarea elementelor de natură tehnică fără depășirea puterii aprobate anterior"),
         ("admin", "modificarea elementelor de natură administrativă fără depășirea puterii aprobate anterior")]


def op_key(*vals) -> str:
    t = " ".join(str(v or "") for v in vals).lower()
    for k, w in (("delgaz", "delgaz"), ("rer", "rețele electrice"), ("rer", "retele electrice"), ("rer", "e-distribu"), ("oltenia", "oltenia"), ("deer", "deer"),
                 ("deer", "distribuție energie electrică"), ("deer", "transilvania"), ("deer", "muntenia nord")):
        if w in t:
            return k
    return ""


def rtype(rec: dict) -> str:
    v = str((rec or {}).get("request_type") or "").strip().casefold()
    return next((k for k, t in TYPES if v in (k.casefold(), t.casefold())), "")


def is_prod(rec: dict) -> bool:
    return "producere" in str(rec.get("kind") or "")


def required(rec: dict) -> list:
    prod = is_prod(rec)
    operator = op_key(rec.get("operator"))
    if operator == "deer":
        case = str(rec.get("deer_case") or "")
        case_norm = case.lower().replace("ț", "t").replace("ș", "s")
        request_norm = str(rec.get("request_type") or "").lower().replace("ț", "t").replace("ș", "s")
        if "spor de putere" in case_norm or "spor" in request_norm:
            case_key = "spor"
        elif "loc nou" in case_norm or "program de finantare" in case_norm:
            case_key = "new"
        elif ">400" in case:
            case_key = "large"
        else:
            case_key = "existing"
        out = ["authenticity", "id", "cf", "prop"]
        if rec.get("applicant_type") == "Persoană fizică":
            out.append("gdpr")
        if rec.get("submission_mode") == "Prin împuternicit":
            out.append("imp")
        if case_key in ("new", "spor"):
            out.append("plan")
        if case_key == "spor":
            out.append("atr_prev")
            if prod:
                out.extend(["pram", "receptie", "schema", "fise"])
        elif case_key == "existing" and prod:
            out.extend(["schema", "fise"])
        elif case_key == "large":
            out.extend(["memoriu", "schema", "fise"])
        if rec.get("applicant_type") == "Persoană juridică" and prod:
            out.append("constatator")
        return list(dict.fromkeys(out))
    request_kind = rtype(rec)
    new = request_kind in ("nou", "temp")
    existing = request_kind in ("spor", "mod_dep", "mod_fara", "admin")
    out = []
    for k, _, req in DOCS:
        if req is True or (req == "prod" and prod) or (k == "cu" and new) or (k == "atr_prev" and existing):
            out.append(k)
    return out


def box(v):
    return "<b>[X]</b>" if v else "[  ]"



DEER_ANNEXES = [
    ("program de finanțare", "Anexa nr. 4", "Cerere de racordare prosumator prin programe de finanțare",
     "https://www.distributie-energie.ro/wp-content/uploads/2023/12/5.-Cerere-de-racordare-prosumator-%E2%80%93-prin-programe-de-finantare-Anexa-4.docx"),
    ("≤10,8", "Anexa nr. 5", "Notificare prosumator / proiect demonstrativ ≤ 10,8 kW",
     "https://www.distributie-energie.ro/wp-content/uploads/2023/12/5.-Loc-consum-si_sau-producere-existent-prevazut_cu_instalatii_de_producere_Anexa-nr-5.docx"),
    (">400", "Procedură dedicată", "Cerere pentru instalații de producere cu putere > 400 kW",
     "https://www.distributie-energie.ro/wp-content/uploads/2024/04/Persoane-fizice-si-juridice-care-au-puterea-instalata-pe-producere-mai-mare-de-400kW.doc"),
    ("actualizare crr", "Anexa nr. 3", "Notificare loc existent — actualizare CRR",
     "https://www.distributie-energie.ro/wp-content/uploads/2023/12/3.-Loc-de-consum-si-producere-existent-actualizare-CRR_Anexa-nr-3.docx"),
    ("spor de putere", "Anexa nr. 1", "Cerere loc existent — spor de putere",
     "https://www.distributie-energie.ro/wp-content/uploads/2023/12/4.-Loc-consum-si_sau-producere-existent-Spor-Putere_Anexa-nr-1.docx"),
    ("loc nou", "Anexa nr. 1", "Cerere loc nou de consum și producere",
     "https://www.distributie-energie.ro/wp-content/uploads/2023/12/1.-Loc-consum-si-producere-nou-_Anexa-nr-1.docx"),
    ("anexa 2", "Anexa nr. 2", "Notificare loc existent",
     "https://www.distributie-energie.ro/wp-content/uploads/2023/12/2.-Loc-de-consum-existent_Anexa-nr-2.docx"),
    ("", "Anexa nr. 2", "Notificare loc existent",
     "https://www.distributie-energie.ro/wp-content/uploads/2023/12/2.-Loc-de-consum-existent_Anexa-nr-2.docx"),
]


def deer_annex_info(rec: dict):
    case = str(rec.get("deer_case") or "").strip().lower()
    if not case:
        return None
    for marker, annex, title, url in DEER_ANNEXES:
        if marker == "" or marker in case:
            return annex, title, url
    return None


def installed_pv_kwp_for_record(rec: dict = None, pros: dict = None):
    """Return DC panel capacity in kWp from confirmed equipment data; never substitute inverter AC kW."""
    rec, pros = rec or {}, pros or {}
    cef = pros.get("cef") or {}
    if cef.get("inverters"):
        try:
            calculated = CEF.check(cef).get("totals", {}).get("kwp")
            if D.num(calculated) > 0:
                return D.num(calculated)
        except Exception:
            pass
        panel_total = 0.0
        complete = True
        for inv in cef.get("inverters") or []:
            count = max(1, int(D.num(inv.get("count")) or 1))
            watts = D.num(inv.get("panel_w"))
            panels = sum(D.num(item.get("panels")) for item in inv.get("strings") or [])
            if panels <= 0 and len(cef.get("inverters") or []) == 1:
                panels = D.num(pros.get("panel_count"))
            if watts <= 0 or panels <= 0:
                complete = False
                break
            panel_total += panels * watts * count
        if complete and panel_total > 0:
            return panel_total / 1000
    flat = CEF.flatten(pros) if pros else {}
    for value in (rec.get("pv_kw"), flat.get("_kwp"), flat.get("panel_kwp"), pros.get("_kwp"), pros.get("panel_kwp")):
        number = D.num(value)
        if number > 0:
            return number
    count = D.num(flat.get("panel_count") or pros.get("panel_count"))
    watts = D.num(flat.get("panel_w") or pros.get("panel_w"))
    if count > 0 and watts > 0:
        return count * watts / 1000
    return None


def deer_annex_issues(rec: dict, pros: dict = None) -> list:
    """Return hard validation issues for DEER cases that have objective eligibility rules."""
    if op_key(rec.get("operator"), (pros or {}).get("operator")) != "deer":
        return []
    case = str(rec.get("deer_case") or "").strip().casefold()
    if not case:
        return ["Selectați procedura DEER aplicabilă înainte de generarea sau depunerea dosarului."]
    issues = []
    request = rtype(rec)
    if "loc nou" in case and request != "nou":
        issues.append("Procedura selectată este Anexa 1 pentru loc nou, dar tipul solicitării nu este «utilizator nou definitiv».")
    elif "spor de putere" in case and request != "spor":
        issues.append("Procedura selectată este Anexa 1 pentru spor de putere, dar tipul solicitării nu este spor de putere.")
    elif "actualizare crr" in case and request != "admin":
        issues.append("Procedura selectată este Anexa 3 pentru actualizare CRR, dar tipul solicitării nu este modificare administrativă.")
    elif "program de finanțare" in case and request != "nou":
        issues.append("Procedura selectată este Anexa 4 pentru program de finanțare, dar tipul solicitării nu este utilizator nou.")
    if "≤10,8" in case or "<=10,8" in case or "<=10.8" in case:
        power = installed_pv_kwp_for_record(rec, pros)
        if power is None:
            issues.append("Pentru Anexa 5 trebuie confirmată puterea instalată de producere (Pi) din fișa tehnică / proiect.")
        elif power > 10.8:
            issues.append(f"Anexa 5 se aplică pentru Pi ≤ 10,8 kW; puterea introdusă este {power:g} kW.")
        supply = str(rec.get("supply") or (pros or {}).get("supply") or ((pros or {}).get("cef") or {}).get("grid", {}).get("supply") or "").casefold()
        if not supply:
            issues.append("Pentru verificarea eligibilității Anexei 5 trebuie selectat tipul branșamentului.")
        elif "trifaz" in supply or "400" in supply:
            issues.append("Anexa 5 selectată este pentru conexiune diferită de trifazic; branșamentul este trifazat.")
    if ">400" in case:
        power = installed_pv_kwp_for_record(rec, pros)
        if power is None:
            issues.append("Procedura dedicată pentru >400 kW necesită puterea instalată de producere.")
        elif power <= 400:
            issues.append(f"Procedura dedicată se aplică pentru putere instalată >400 kW; puterea introdusă este {power:g} kW.")
    return issues


def storage_details(pros: dict, rec: dict) -> list:
    cef = (pros or {}).get("cef") or {}
    details = []
    for inv in cef.get("inverters") or []:
        if inv.get("battery_mode") == "own":
            battery = inv.get("battery") or {}
            if battery:
                quantity = max(1, int(D.num(inv.get("count")) or 1))
                modules = max(1, int(D.num(battery.get("modules")) or 1))
                details.append({**battery, "modules": modules * quantity})
    details.extend(cef.get("shared_batteries") or [])
    if not details and ((rec or {}).get("battery") or (rec or {}).get("storage")):
        value = (rec or {}).get("battery") or (rec or {}).get("storage")
        details.append({
            "model": value if isinstance(value, str) else value.get("model"),
            "modules": (rec or {}).get("battery_count"),
            "kwh": (rec or {}).get("battery_kwh"),
            "ah": (rec or {}).get("battery_ah"),
            "voltage": (rec or {}).get("battery_v"),
            "i_max": (rec or {}).get("battery_i_max"),
            "coupling": (rec or {}).get("battery_coupling"),
        })
    return details


def storage_row(battery: dict) -> list:
    modules = battery.get("modules")
    kwh = battery.get("kwh")
    if kwh not in (None, "") and modules not in (None, ""):
        try:
            capacity = f"{D.fmt(float(str(kwh).replace(',', '.')) * float(str(modules).replace(',', '.')), 2)} kWh total ({modules} × {kwh} kWh)"
        except (TypeError, ValueError):
            capacity = f"{kwh} kWh"
    else:
        capacity = f"{kwh} kWh" if kwh not in (None, "") else M
    coupling = battery.get("coupling")
    return [escape(str(battery.get("model") or M)), escape(str(modules or M)), escape(capacity),
            escape(str(battery.get("ah") or M)), escape(str(battery.get("voltage") or M)),
            escape(str(battery.get("i_max") or M)), escape(str(coupling).upper() if coupling else "Nespecificată")]


def render_deer_annex(rec: dict, client: dict, settings: dict, pros: dict, attached: set) -> bytes:
    """Create a clearly-labeled work prefill sheet for the selected DEER annex."""
    info = deer_annex_info(rec)
    if not info or op_key(rec.get("operator"), (pros or {}).get("operator"), (client or {}).get("operator")) != "deer":
        return b""
    annex, title, url = info
    c = client or {}
    p = CEF.flatten(pros) if pros else {}
    pv_kwp = installed_pv_kwp_for_record(rec, pros)
    g = lambda d, k: escape(str(d.get(k))) if d.get(k) not in (None, "") else M
    P = lambda t, k="p": Paragraph(t, D.S[k])
    T = lambda rows, w: D.pdf_table(rows, w, True, 8, 178)
    address = rec.get("address") or c.get("address")
    pod = rec.get("pod") or c.get("pod")
    cf = rec.get("cf_nr") or c.get("cf_nr")
    cadastral = rec.get("cadastral_nr") or c.get("cadastral_nr")
    story = [
        P(f"{annex.upper()} — {escape(title)}", "h"),
        P("FIȘĂ DE PRECOMPLETARE PENTRU DOSARUL ATR", "sub"),
        P("Document de lucru generat din datele salvate în DNS Energy. Nu este formularul oficial DEER și trebuie verificat și transcris în formularul curent al operatorului înainte de depunere.", "small"),
        Spacer(1, 4),
        P("1. Solicitant și reprezentare", "b"),
        T([["Câmp", "Date din client / dosar"],
           ["Solicitant", g(c, "name")], ["Categorie solicitant", g(rec, "applicant_type")],
           ["CNP / CUI", g(c, "idcode")], ["Telefon", g(c, "phone")], ["E-mail", g(c, "email")],
           ["Reprezentant", g(c, "rep_name")], ["Modalitate depunere", g(rec, "submission_mode")]], [52, 126]),
        Spacer(1, 4),
        P("2. Loc de consum / obiectiv", "b"),
        T([["Câmp", "Date din cerere"],
           ["Denumire obiectiv", g(rec, "object_name")], ["Adresă", g({"address": address}, "address")],
           ["POD", g({"pod": pod}, "pod")], ["Carte funciară", g({"cf": cf}, "cf")],
           ["Număr cadastral", g({"cadastral": cadastral}, "cadastral")],
           ["Operator", g(rec, "operator")], ["Tip loc", g(rec, "kind")]], [52, 126]),
        Spacer(1, 4),
        P("3. Cerere și puteri solicitate", "b"),
        T([["Câmp", "Date din cerere"],
           ["Procedură DEER selectată", escape(str(rec.get("deer_case") or annex))],
           ["Tip solicitare", g(rec, "request_type")], ["Tip racord", g(rec, "supply")],
           ["Putere absorbită", f"{g(rec, 'p_abs_kw')} kW / {g(rec, 'p_abs_kva')} kVA"],
           ["Putere instalată producere (DC)", f"{g({'pv_kwp': pv_kwp}, 'pv_kwp')} kWp"],
           ["Putere maximă evacuată", f"{g(rec, 'p_evac_kw')} kW / {g(rec, 'p_evac_kva')} kVA"],
           ["Descriere modificări", g(rec, "current_supply_desc")]], [52, 126]),
        Spacer(1, 4),
        P("4. Echipamente din dosarul tehnic asociat", "b"),
    ]
    cef = (pros or {}).get("cef") or {}
    invs = cef.get("inverters") or []
    equipment = [["Invertor", "Putere", "Panou / număr / Wp", "Stringuri"]]
    for inv in invs:
        strings = inv.get("strings") or []
        quantity = max(1, int(D.num(inv.get("count")) or 1))
        panel_count = sum(int(D.num(x.get("panels"))) for x in strings) * quantity
        if not panel_count and inv is invs[0]:
            panel_count = int(D.num(p.get("panel_count")) or 0)
        panel_w = D.num(inv.get("panel_w"))
        string_text = ", ".join(str(int(D.num(x.get("panels")))) for x in strings if D.num(x.get("panels")))
        equipment.append([escape(str(inv.get("brand_model") or M)), escape(str(inv.get("p_nom_kw") or M)) + " kW",
                          escape(str(inv.get("panel_model") or M)) + (f" / {panel_count} × {panel_w:g} Wp" if panel_count and panel_w else ""),
                          escape(string_text or M)])
    if not invs:
        equipment.append([M, M, M, M])
    story.append(T(equipment, [50, 22, 70, 36]))
    batteries = storage_details(pros, rec)
    story.append(P("Sisteme de stocare: " + ("DA" if batteries else ("NU" if invs else "nespecificat")), "small"))
    if batteries:
        story.append(T([["Model", "Buc.", "Capacitate", "Ah", "V", "Cuplare"]] + [storage_row(b) for b in batteries], [46, 16, 46, 18, 20, 30]))
    story.append(P("5. Documente anexate pentru această procedură", "b"))
    rows = [["Document solicitat", "Stare în dosar"]]
    for key, label, _ in DOCS:
        if key in required(rec):
            rows.append([escape(label), "Atașat" if key in attached else "Lipsește"])
    story += [T(rows, [138, 40]), Spacer(1, 5),
              P("Formularul oficial selectat trebuie descărcat și verificat la sursa DEER:", "small"),
              P(f'<link href="{escape(url)}" color="blue">{escape(url)}</link>', "small"),
              Spacer(1, 5),
              P("Câmpurile care lipsesc din datele clientului sau din cerere rămân necompletate aici; nu se estimează și nu se inventează valori.", "small")]
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=14 * mm, bottomMargin=14 * mm,
                            title=f"Precompletare {annex} — {title}")
    doc.build(story)
    return buf.getvalue()


def _deer_docx_template(rec: dict) -> str:
    case = str((rec or {}).get("deer_case") or "").strip().lower()
    if "program de finanțare" in case or "program de finantare" in case:
        return "anexa4_funding.docx.b64"
    if "≤10,8" in case or "≤ 10,8" in case or "<=10,8" in case:
        return "anexa5_simple.docx.b64"
    if "actualizare crr" in case or "anexa nr. 3" in case:
        return "anexa3_update.docx.b64"
    if "spor" in case:
        return "anexa1_spor.docx.b64"
    if "loc nou" in case:
        return "anexa1_new.docx.b64"
    if ">400" in case:
        return None
    if "anexa 2" in case or "existent" in case:
        return "anexa2_existing.docx.b64"
    return "anexa2_existing.docx.b64"


def _deer_value(value):
    if value in (None, ""):
        return ""
    return str(value).strip()


def _deer_cell(cell, value):
    value = _deer_value(value)
    if value:
        cell.text = value


def _deer_table_row(table, index, values):
    while len(table.rows) <= index:
        table.add_row()
    for col, value in enumerate(values):
        if col < len(table.columns):
            _deer_cell(table.cell(index, col), value)


def render_deer_official_docx(rec: dict, client: dict, settings: dict, pros: dict) -> bytes:
    """Fill the official DEER DOCX template with known dossier values; never infer missing facts."""
    from base64 import b64decode
    from pathlib import Path
    from docx import Document

    if op_key((rec or {}).get("operator"), (pros or {}).get("operator"), (client or {}).get("operator")) != "deer":
        raise ValueError("Precompletarea formularului oficial este disponibilă doar pentru DEER.")
    filename = _deer_docx_template(rec)
    if not filename:
        raise ValueError("Procedura DEER pentru instalații de peste 400 kW este disponibilă pe site-ul operatorului în format DOC; descărcați formularul oficial de acolo.")
    template_path = Path(__file__).parent / "templates" / "deer" / filename
    try:
        raw = b64decode(template_path.read_text(encoding="ascii"), validate=True)
        doc = Document(io.BytesIO(raw))
    except Exception as exc:
        raise RuntimeError(f"Șablonul oficial DEER nu a putut fi încărcat: {exc}") from exc

    c = client or {}
    rec = rec or {}
    pros = pros or {}
    cef = pros.get("cef") or {}
    inverters = cef.get("inverters") or []
    batteries = storage_details(pros, rec)
    case = str(rec.get("deer_case") or "").lower().replace("ț", "t").replace("ș", "s")
    is_new = "loc nou" in case or "program de finantare" in case
    # Applicant details are substituted only where the official form has an unambiguous field.
    import re
    applicant = next((p for p in doc.paragraphs if "Utilizatorul" in p.text), None)
    if applicant:
        fields = {
            0: c.get("name"),
            1: c.get("county"),
            2: c.get("locality"),
            10: c.get("phone"),
            13: c.get("email"),
            14: c.get("cui") or (c.get("idcode") if str(rec.get("applicant_type") or "").lower() in ("pj", "persoană juridică", "persoana juridica") else None),
            15: c.get("idcode") if str(rec.get("applicant_type") or "").lower() not in ("pj", "persoană juridică", "persoana juridica") else None,
            16: c.get("registration_no"),
            17: c.get("rep_name"),
            18: c.get("rep_role"),
        }
        counter = 0
        def replace_field(match):
            nonlocal counter
            value = fields.get(counter)
            counter += 1
            return _deer_value(value) or match.group(0)
        original = applicant.text
        applicant.clear()
        applicant.add_run(re.sub(r"\.{4,}", replace_field, original))

    # Select the request type without touching unrelated declaration checkboxes.
    for paragraph in doc.paragraphs:
        if "Loc de consum şi de producere nou" in paragraph.text and "existent" not in paragraph.text.lower():
            paragraph.text = paragraph.text.replace("[]", "[X]" if is_new else "[]", 1)
        elif "Loc de consum existent/Loc de consum şi de producere existent" in paragraph.text:
            paragraph.text = paragraph.text.replace("[]", "[X]" if not is_new else "[]", 1)

    # Official DEER Anexa 1 and 4/5 technical tables share this ordering; Anexa 2 has fewer tables.
    tables = doc.tables
    is_annex2 = filename == "anexa2_existing.docx.b64"
    power_table = 0
    if not is_annex2 and filename != "anexa3_update.docx.b64" and len(tables) > power_table:
        _deer_cell(tables[power_table].cell(0, 1), rec.get("p_abs_kw") or pros.get("bransament_kw"))
        _deer_cell(tables[power_table].cell(1, 1), rec.get("p_abs_kva"))

    if is_annex2:
        export_table, panel_table, inverter_table, detail_table = 0, 2, 3, 5
    else:
        export_table, panel_table, inverter_table, detail_table = 2, 5, 6, 8
    if len(tables) > export_table and not is_annex2 and filename != "anexa3_update.docx.b64":
        _deer_cell(tables[export_table].cell(0, 1), first_present(rec.get("p_evac_kw"), pros.get("power")))
        _deer_cell(tables[export_table].cell(1, 1), rec.get("p_evac_kva"))
    elif is_annex2 and len(tables) > 0:
        # Anexa 2 does not have an export power table; its power is printed in the request sentence.
        pass

    if len(tables) > panel_table:
        table = tables[panel_table]
        row = 2
        for idx, inv in enumerate(inverters, 1):
            strings = inv.get("strings") or []
            try:
                panels = sum(int(D.num(x.get("panels"))) for x in strings) * max(1, int(D.num(inv.get("count")) or 1))
            except (TypeError, ValueError):
                panels = 0
            if not panels and idx == 1:
                panels = int(D.num(rec.get("panel_count") or pros.get("panel_count") or CEF.flatten(pros).get("panel_count")) or 0)
            watts = D.num(inv.get("panel_w"))
            model = inv.get("panel_model")
            total_kw = panels * watts / 1000 if panels and watts else None
            _deer_table_row(table, row, [str(idx), panels or "", model,
                (f"{watts / 1000:g}" if watts else ""), (f"{total_kw:g}" if total_kw else ""),
                "", (f"{total_kw:g}" if total_kw else ""), ""])
            row += 1

    if len(tables) > inverter_table:
        table = tables[inverter_table]
        row = 2
        for idx, inv in enumerate(inverters, 1):
            model = inv.get("brand_model")
            rated = D.num(inv.get("p_nom_kw"))
            unit_limit = D.num(inv.get("p_limit_kw"))
            count = max(1, int(D.num(inv.get("count")) or 1))
            limit = f"{unit_limit:g}" if unit_limit is not None and unit_limit > 0 else ""
            plant_max = rated * count if rated is not None and rated > 0 else ""
            note = f"Limitare evacuare configurată: {unit_limit:g} kW/invertor" if unit_limit is not None and unit_limit > 0 else ""
            if not is_annex2 and inv.get("battery_mode") == "own" and inv.get("battery", {}).get("model"):
                battery = inv["battery"]
                details = []
                for key, label in (("kwh", "kWh"), ("voltage", "V"), ("ah", "Ah"), ("i_max", "I max/modul A")):
                    value = D.num(battery.get(key))
                    if value is not None and value > 0:
                        details.append(f"{value:g} {label}")
                if battery.get("coupling"):
                    details.append("cuplare " + str(battery["coupling"]))
                description = "Baterie: " + str(battery["model"])
                if battery.get("modules"):
                    description += " × " + str(battery["modules"])
                if details:
                    description += " (" + ", ".join(details) + ")"
                note = (note + "; " if note else "") + description
            _deer_table_row(table, row, [str(idx), count, model, "",
                f"{rated:g}" if rated is not None and rated > 0 else "", limit,
                f"{plant_max:g}" if plant_max else "", note])
            row += 1
        if inverters and len(table.rows) > row:
            total = sum(D.num(inv.get("p_nom_kw")) * D.num(inv.get("count") or 1) for inv in inverters)
            _deer_cell(table.cell(len(table.rows) - 1, 4), f"{total:g}" if total else "")
        # List batteries as equipment rows, matching inverter visibility in Annex 2.
        # Do not overwrite the template's total row; append battery rows after it.
        for battery_idx, battery in enumerate(batteries, 1):
            modules = max(1, int(D.num(battery.get("modules")) or 1))
            specs = []
            for key, label, unit in (("voltage", "U", " V"), ("ah", "Capacitate", " Ah"),
                                     ("i_max", "I max", " A"), ("coupling", "Cuplare", "")):
                value = battery.get(key)
                if value not in (None, ""):
                    specs.append(f"{label}: {value}{unit}")
            capacity = D.num(battery.get("kwh"))
            if capacity:
                capacity *= modules
            battery_row = len(table.rows)
            _deer_table_row(table, battery_row, [f"B{battery_idx}", modules, battery.get("model") or "",
                "Baterie", f"{capacity:g}" if capacity else "", "", "", "; ".join(specs)])

    if len(tables) > detail_table:
        detail = tables[detail_table]
        supply_details = []
        for inv in inverters:
            if inv.get("brand_model"):
                supply_details.append("Invertor: " + str(inv["brand_model"]))
        for battery in ([] if is_annex2 else batteries):
            desc = "Baterie: " + str(battery.get("model") or "model necompletat")
            if battery.get("modules"):
                desc += " × " + str(battery["modules"])
            if battery.get("coupling"):
                desc += " (" + str(battery["coupling"]) + ")"
            if D.num(battery.get("i_max")) is not None:
                desc += f"; I max/modul {D.num(battery.get('i_max')):g} A"
            supply_details.append(desc)
        if supply_details:
            _deer_cell(detail.cell(0, 1), "; ".join(supply_details))
        amp_hours = [] if is_annex2 else [str(b.get("ah")) for b in batteries if b.get("ah") not in (None, "")]
        if amp_hours:
            _deer_cell(detail.cell(1, 1), ", ".join(amp_hours))

    # Populate the official location line for Anexa 1 variants using only stored values.
    if not is_annex2:
        location_request = next((p for p in doc.paragraphs if "obţinerea avizului tehnic de racordare" in p.text.lower()), None)
        if location_request:
            import re
            address_fields = {
                0: c.get("county"), 1: c.get("locality"),
                2: rec.get("address") or c.get("address"),
                8: rec.get("cf_nr") or c.get("cf_nr"),
                9: rec.get("cadastral_nr") or c.get("cadastral_nr"),
            }
            address_count = 0
            def replace_address_field(match):
                nonlocal address_count
                value = address_fields.get(address_count)
                address_count += 1
                return _deer_value(value) or match.group(0)
            original = location_request.text
            location_request.clear()
            location_request.add_run(re.sub(r"\.{4,}", replace_address_field, original))

    # Anexa 2 puts client/location/POD/export data in its request sentence.
    if is_annex2:
        request = next((p for p in doc.paragraphs if "notific prin prezenta racordarea" in p.text.lower()), None)
        if request:
            # Replace only explicit fields; never split or reinterpret a free-form address.
            import re
            field_values = {
                0: c.get("county"), 1: c.get("locality"),
                8: rec.get("cf_nr") or c.get("cf_nr"),
                9: rec.get("cadastral_nr") or c.get("cadastral_nr"),
                10: rec.get("pod") or c.get("pod"),
                11: first_present(rec.get("p_evac_kw"), pros.get("power")),
                12: rec.get("p_evac_kva"),
            }
            count = 0
            def replace_request_field(match):
                nonlocal count
                value = field_values.get(count)
                count += 1
                return _deer_value(value) or match.group(0)
            original = request.text
            request.clear()
            request.add_run(re.sub(r"\.{4,}", replace_request_field, original))

    doc.add_paragraph("Precompletat automat din datele salvate în dosar. Câmpurile fără date au rămas necompletate; verificați formularul și semnați-l înainte de depunere.")
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def render_atr(rec: dict, client: dict, s: dict, pros: dict, attached: set) -> bytes:
    st, c = D.S, client or {}
    p = CEF.flatten(pros) if pros else {}
    g = lambda d, k: escape(str(d.get(k))) if d.get(k) not in (None, "") else M
    P = lambda t, k="p": Paragraph(t, st[k])
    T = lambda rows, w: D.pdf_table(rows, w, True, 7.6, 178)
    prod = is_prod(rec)
    op = op_key(rec.get("operator"), (pros or {}).get("operator"), c.get("operator"))
    rt = rtype(rec)
    deer_case = str(rec.get("deer_case") or "").strip()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=14 * mm, bottomMargin=14 * mm, title="Cerere ATR")
    story = [P(f"Către: <b>{escape(OP_NAMES[op])}</b>" + (f" — {g(rec, 'operator')}" if rec.get("operator") else ""), "small"),
             P("Nr. înregistrare ____________ / Data ____________", "small"),
             P("CERERE ȘI CHESTIONAR", "h"),
             P("pentru obținerea avizului tehnic de racordare — " + ("loc de consum și de producere (prosumator)" if prod else "loc de consum"), "sub"),
             P("Tip solicitare:<br/>" + "<br/>".join(f"{box(rt == k)} {t}" for k, t in TYPES)),
             *([P("<b>ATENȚIE:</b> selectați tipul solicitării înainte de depunere; cererea nu poate fi clasificată automat din datele prosumatorului.", "small")] if not rt else []),
             *([P("Formular DEER selectat — precompletare cu datele clientului", "b"),
                P(escape(deer_case), "sub"),
                T([["Date preluate din dosarul clientului", "Valoare"],
                   ["Solicitant", g(c, "name")], ["Persoană", g(rec, "applicant_type")],
                   ["CNP / CUI", g(c, "idcode")], ["Telefon / e-mail", f"{g(c, 'phone')} / {g(c, 'email')}"],
                   ["Adresă loc de consum", f"{g(c, 'county')}, {g(c, 'locality')}, {g(rec, 'address') if rec.get('address') else g(c, 'address')}"],
                   ["POD", g(rec, "pod") if rec.get("pod") else g(c, "pod")],
                   ["CF / cadastral", f"{g(rec, 'cf_nr') if rec.get('cf_nr') else g(c, 'cf_nr')} / {g(rec, 'cadastral_nr') if rec.get('cadastral_nr') else g(c, 'cadastral_nr')}"],
                   ["Putere absorbită solicitată", f"{g(rec, 'p_abs_kw')} kW / {g(rec, 'p_abs_kva')} kVA"],
                   ["Tip racord", g(rec, "supply")]], [54, 124]),
                P("Câmpurile fără date în dosar rămân necompletate pentru verificare. Acest PDF este o precompletare de lucru și trebuie confruntat cu formularul DEER curent înainte de depunere.", "small")]
                if deer_case else []),
             P("1. Date de identificare ale solicitantului și ale consultantului de specialitate", "b"),
             P(f"1.1 Numele solicitantului: <b>{g(c, 'name')}</b>; adresa: județ {g(c, 'county')}, localitatea {g(c, 'locality')}, {g(c, 'address')}; telefon {g(c, 'phone')}, "
               f"e-mail {g(c, 'email')}; nr. O.R.C. {g(c, 'regcom') if c.get('regcom') else '—'}; cod fiscal / CNP {g(c, 'idcode')}; cont {g(c, 'iban')} deschis la {g(c, 'bank')}; "
               f"reprezentat prin {g(c, 'rep_name') if c.get('rep_name') else '—'} în calitate de {g(c, 'rep_role') if c.get('rep_role') else '—'}."),
             P(f"1.2 Consultant de specialitate / împuternicit: <b>{g(s, 'company')}</b>, telefon {g(s, 'phone')}, e-mail {g(s, 'email')}, "
               f"atestat ANRE nr./data {g(s, 'anre_attest')}{(' / ' + escape(str(s.get('anre_issued')))) if s.get('anre_issued') else ''}."),
             P("2. Date generale asupra obiectivului pentru care se solicită avizul tehnic de racordare", "b"),
             P(f"2.1 POD loc de consum: {g(rec, 'pod') if rec.get('pod') else g(c, 'pod')}"),
             P(f"2.2 Denumirea obiectivului: {g(rec, 'object_name')}<br/>2.3 Adresa obiectivului: județ {g(c, 'county')}, localitatea {g(c, 'locality')}, "
               f"{g(rec, 'address') if rec.get('address') else g(c, 'address')}; CF nr. {g(rec, 'cf_nr') if rec.get('cf_nr') else g(c, 'cf_nr')}, nr. cadastral "
               f"{g(rec, 'cadastral_nr') if rec.get('cadastral_nr') else g(c, 'cadastral_nr')}"),
             P(f"2.3 Specificul activității: {g(rec, 'activity')}; număr de schimburi {g(rec, 'shifts')}; zile lucrătoare / săptămână {g(rec, 'workdays')} (doar persoane juridice)"),
             P(f"2.4 Modul de alimentare actual și descrierea modificărilor de natură tehnică (spor de putere): {g(rec, 'current_supply_desc') if rt not in ('nou', 'temp') else '—'}"),
             P("3. Date energetice ale locului de consum", "b"),
             P(f"3.1 Tipul de racord solicitat: {box(rec.get('supply') == 'Monofazat')} monofazat {box(rec.get('supply') == 'Trifazat')} trifazat"),
             P("3.2 Lista receptoarelor (puterea și tensiunea nominală; regim cu șocuri, deformant, dezechilibrat unde este cazul):")]
    recs = rec.get("receptors") or []
    story.append(T([["Nr. crt.", "Denumirea receptoarelor", "Pi (kW)", "Observații (caracteristici, regim de funcționare)"]] +
                   [[str(i + 1), escape(r.get("name") or ""), escape(str(r.get("kw") or "")), escape(r.get("notes") or "")] for i, r in enumerate(recs)] +
                   ([["", "", "", ""]] * (0 if recs else 3)) +
                   [["", "Total Pi (kW)", D.fmt(sum(D.num(r.get("kw")) for r in recs)) if recs else M, ""],
                    ["", "Puterea maximă simultan absorbită", f"{g(rec, 'p_abs_kw')} kW", f"{g(rec, 'p_abs_kva')} kVA"]], [8, 42, 16, 34]))
    story += [P(f"3.3 Puterea celui mai mare motor: {g(rec, 'motor_kw')} kW; regim de lucru {g(rec, 'motor_regime')}; mod de pornire {g(rec, 'motor_start')}"),
              P(f"3.4 Durata maximă de restabilire a alimentării acceptată: {g(rec, 'restore_h')} ore; număr căi de alimentare solicitate: {g(rec, 'paths')}"),
              P(f"3.5 Factorul de putere mediu: cos φ = {g(rec, 'cos_phi')}<br/>3.6 Cerințe privind grupul de măsurare: {g(rec, 'metering')}"),
              P(f"3.7 Data estimată pentru punerea sub tensiune finală: {escape(D.ro_date(rec.get('energize_date')) or M)}; 3.8 etape: {g(rec, 'stages') if rec.get('stages') else '—'}")]
    if prod:
        invs = (pros.get("cef") or {}).get("inverters") or [] if pros else []
        pv_kw = installed_pv_kwp_for_record(rec, pros)
        story += [P("4. Date tehnice și energetice ale locului de producere (centrală fotovoltaică)", "b"),
                  T([["Nr.", "Date energetice ale producătorului", "Simbol", "U.M.", "Situația existentă", "Situația finală"],
                     ["1", "Puterea totală instalată (panouri c.c.)", "Pi", "kWp", escape(str(rec.get("pv_existing_kw") or "0")), escape(str(pv_kw or M))],
                     ["2", "Puterea maximă simultan evacuată în rețea", "Pa", "kW", "", escape(str(first_present(rec.get("p_evac_kw"), p.get("power"), M)))],
                     ["", "", "Sa", "kVA", "", escape(str(first_present(rec.get("p_evac_kva"), p.get("apparent_kva"), M)))],
                     ["3", "Consumul propriu al centralei (servicii interne)", "Pi / Pa", "kW", "", f"{escape(str(p.get('internal_pi') or rec.get('internal_kw') or M))}"],
                     ["4", "Factorul de putere", "cos φ", "", "", escape(str(rec.get("cos_phi_prod") or "1"))]], [6, 42, 10, 8, 16, 18]),
                  Spacer(1, 3), P("Invertoare / unități generatoare:")]
        rows = [["Nr.", "Tip invertor", "Nr. buc.", "Pn (kW c.a.)", "Panouri (nr. × Wp)", "Pi c.c. (kWp)"]]
        for i, inv in enumerate(invs):
            quantity = max(1, int(D.num(inv.get("count") or inv.get("quantity")) or 1))
            if any(sum(int(D.num(x.get("panels"))) for x in entry.get("strings") or []) for entry in invs):
                n = sum(int(D.num(x.get("panels"))) for x in inv.get("strings") or []) * quantity
            else:
                n = D.num(p.get("panel_count")) if i == 0 else 0
            w = D.num(inv.get("panel_w"))
            rows.append([str(i + 1), escape(inv.get("brand_model") or ""), escape(str(quantity)), escape(str(inv.get("p_nom_kw") or "")), f"{n:g} × {w:g}" if n else "", D.fmt(n * w / 1000, 2) if n and w else ""])
        if not invs:
            rows.append([""] * 6)
        batteries = storage_details(pros, rec)
        configured = bool((pros or {}).get("cef", {}).get("inverters") or (pros or {}).get("inverter"))
        storage_state = f"{box(True)} DA" if batteries else (f"{box(True)} NU" if configured else "nespecificat — confirmați prezența bateriei")
        story += [T(rows, [6, 40, 10, 14, 18, 12]),
                  P(f"Sisteme de stocare: {storage_state}"),
                  *([T([["Model", "Buc.", "Capacitate", "Ah", "V", "Cuplare"]] + [storage_row(b) for b in batteries], [46, 16, 46, 18, 20, 30])] if batteries else []),
                  P(f"Data prevăzută pentru PIF a centralei: {escape(D.ro_date(rec.get('pif_date')) or M)}")]
    n = 5 if prod else 4
    story += [P(f"{n}. Alte informații", "b"), P(g(rec, "other") if rec.get("other") else M),
              P(f"{n + 1}. Documente anexate (copii)", "b")]
    req = set(required(rec))
    for k, t, _ in DOCS:
        if k in req or k in attached:
            story.append(P(f"{box(k in attached)} {t}", "li"))
    story.append(P(f"{box(True)} Declarație pe propria răspundere privind autenticitatea datelor și conformitatea cu originalul a documentelor anexate", "li"))
    story += [Spacer(1, 6), P("Declar pe propria răspundere că datele și informațiile cuprinse în prezenta cerere sunt autentice și că documentele anexate, în copie, sunt conforme cu originalul. "
                               "Sunt de acord cu prelucrarea datelor cu caracter personal în scopul soluționării cererii de racordare."),
              D.pdf_sign([f"Data: {escape(D.ro_date(rec.get('date')) or M)}"], ["Solicitant / Împuternicit,", f"<b>{g(c, 'name')}</b> / <b>{g(s, 'company')}</b>",
                                                                                    escape(s.get("rep") or "") or M, "(numele, prenumele și semnătura)", "(ștampilă)"], s),
              Spacer(1, 4), P(f"Model întocmit după cererea-tip a operatorului ({escape(OP_NAMES[op])}) și Regulamentul de racordare (Ord. ANRE 59/2013). "
                              "Verificați în Registrul formularelor oficiale versiunea curentă publicată de operator.", "small")]
    doc.build(story)
    return buf.getvalue()



def render_operator_annex_prep(rec: dict, client: dict, pros: dict, operator_key: str, annex_no: int, form_title: str = "") -> bytes:
    """Render a data worksheet to help complete the selected operator's official annex.

    This is explicitly a work sheet, never represented as the operator's official form.
    Only values present in the client/ATR/prosumer records are printed.
    """
    rec, client, pros = rec or {}, client or {}, pros or {}
    cef = pros.get("cef") or {}
    inverters = cef.get("inverters") or []
    p = CEF.flatten(pros) if pros else {}
    installed_pv_kwp = installed_pv_kwp_for_record(rec, pros)
    limit_values = []
    for inv in inverters:
        model = str(inv.get("brand_model") or "Invertor")
        unit_count = max(1, int(D.num(inv.get("count")) or 1))
        unit_limit = D.num(inv.get("p_limit_kw"))
        if unit_limit is None or unit_limit <= 0:
            limit_values.append(f"{model}: limită nedeclarată")
        else:
            limit_values.append(f"{model}: {unit_limit:g} kW/buc. × {unit_count} buc.")
    inverter_limit = "; ".join(limit_values) if limit_values else "Nedeclarată în configurarea invertorului"
    applicant_type = str(rec.get("applicant_type") or "").casefold()
    is_company = "jurid" in applicant_type or applicant_type == "pj"
    applicant_identifier = (client.get("cui") or client.get("idcode")) if is_company else client.get("idcode")
    applicant_address = (client.get("hq_address") or client.get("address")) if is_company else (client.get("address") or client.get("hq_address"))
    v = lambda value: escape(str(value)) if value not in (None, "") else M
    P = lambda text, style="p": Paragraph(text, D.S[style])
    T = lambda rows, widths: D.pdf_table(rows, widths, True, 7.6, 178)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=14 * mm, bottomMargin=14 * mm,
                            title=f"Fișă precompletare Anexa {annex_no}")
    operator = OP_NAMES.get(operator_key, rec.get("operator") or "Operatorul de distribuție")
    story = [
        P(f"FIȘĂ DE PRECOMPLETARE — ANEXA {int(annex_no)}", "h"),
        P(escape(operator), "sub"),
        P("DOCUMENT DE LUCRU — NU ESTE FORMULARUL OFICIAL AL OPERATORULUI. Folosiți datele pentru completarea formularului oficial atașat și verificați fiecare câmp înainte de depunere. Valorile lipsă nu sunt deduse sau inventate.", "small"),
        Spacer(1, 5),
    ]
    if form_title:
        story.append(P(f"Formular din registru: {escape(form_title)}", "small"))
    story += [
        P("1. Solicitant", "b"),
        T([["Câmp", "Date disponibile"], ["Nume / denumire", v(client.get("name") or rec.get("client_name"))],
           ["Tip solicitant", v(rec.get("applicant_type"))], ["CNP / CUI", v(applicant_identifier)],
           ["Domiciliu / sediu din act", v(applicant_address)],
           ["Telefon", v(client.get("phone"))], ["E-mail", v(client.get("email"))],
           ["Reprezentant", v(client.get("rep_name"))]], [48, 130]),
        Spacer(1, 4),
        P("2. Loc de consum / racordare", "b"),
        T([["Câmp", "Date disponibile"], ["Adresă loc de consum", v(rec.get("address") or client.get("address"))],
           ["Localitate / județ", v(rec.get("locality") or client.get("locality")) + " / " + v(rec.get("county") or client.get("county"))],
           ["POD", v(rec.get("pod") or client.get("pod"))], ["Nr. carte funciară", v(rec.get("cf_nr") or client.get("cf_nr"))],
           ["Nr. cadastral", v(rec.get("cadastral_nr") or client.get("cadastral_nr"))],
           ["Operator", v(rec.get("operator") or operator)], ["Procedură", v(rec.get("deer_case") or rec.get("request_type"))],
           ["Tip branșament", v(rec.get("supply") or pros.get("supply"))]], [48, 130]),
        Spacer(1, 4),
        P("3. Puteri și parametri solicitați", "b"),
        T([["Câmp", "Valoare disponibilă"], ["Putere absorbită solicitată", v(rec.get("p_abs_kw")) + " kW / " + v(rec.get("p_abs_kva")) + " kVA"],
           ["Putere instalată fotovoltaică (DC)", v(installed_pv_kwp) + " kWp"],
           ["Putere evacuată solicitată", v(first_present(rec.get("p_evac_kw"), pros.get("power"))) + " kW / " + v(rec.get("p_evac_kva")) + " kVA"],
           ["Limită de injecție declarată în setările invertorului", v(inverter_limit)]], [48, 130]),
        Spacer(1, 4),
        P("4. Echipamente declarate", "b"),
    ]
    pv_rows = [["Invertor", "Buc.", "Putere (kW)", "Panouri / stringuri"]]
    for inv in inverters:
        count = max(1, int(D.num(inv.get("count")) or 1))
        strings = inv.get("strings") or []
        panels = sum(int(D.num(x.get("panels"))) for x in strings) * count
        string_text = ", ".join(str(int(D.num(x.get("panels")))) for x in strings if D.num(x.get("panels")))
        panel = str(inv.get("panel_model") or p.get("panels") or "")
        watts = D.num(inv.get("panel_w") or p.get("panel_w"))
        pv_rows.append([v(inv.get("brand_model")), str(count), v(inv.get("p_nom_kw")),
                        v(panel) + (f" — {panels} × {watts:g} Wp" if panels and watts else "") + (f"; șiruri {string_text}" if string_text else "")])
    if not inverters and (pros.get("inverter") or rec.get("inverter")):
        pv_rows.append([v(pros.get("inverter") or rec.get("inverter")), v(pros.get("inverter_count") or 1),
                        v(pros.get("inverter_kw") or rec.get("inverter_kw")), v(pros.get("panels") or rec.get("panels"))])
    if len(pv_rows) == 1:
        pv_rows.append([M, M, M, M])
    story.append(T(pv_rows, [54, 12, 24, 88]))
    batteries = storage_details(pros, rec)
    story.append(P("Stocare: " + ("DA" if batteries else ("NU" if inverters else "nespecificat")), "small"))
    if batteries:
        story.append(T([["Model baterie", "Module", "Capacitate", "Ah", "V", "I max (A)", "Cuplare"]] + [storage_row(x) for x in batteries], [42, 12, 38, 15, 15, 20, 26]))
    story += [Spacer(1, 4), P("5. Verificare înainte de depunere", "b"),
              P("Comparați această fișă cu formularul oficial atașat. Confirmați datele de identificare, POD, tipul cererii, puterile aprobate/solicitate și cerințele specifice operatorului. Fișa nu înlocuiește semnătura solicitantului, documentele justificative sau verificarea proiectantului.", "small")]
    doc.build(story)
    return buf.getvalue()
