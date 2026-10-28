from datetime import datetime, timezone, date, timedelta
import asyncio
import hmac
import re
import os
import uuid
from html import escape as hesc

from fastapi import HTTPException, Request, Depends, UploadFile, File, Header, Query, Response, BackgroundTasks

import server as S
from mailer import send_email

db, api, get_current_user = S.db, S.api_router, S.get_current_user
TYPES = ["A1", "A2", "A3", "B", "Bi", "Bp", "C1A", "C1B", "C2A", "C2B", "D1", "D2", "E1", "E2"]
STATUSES = ["în așteptare", "aprobat", "respins", "expirat"]
FIELDS = ("company", "cui", "regcom", "contact_name", "contact_email", "contact_phone", "attest_no", "attest_types", "attest_other", "attest_valid", "holder_name", "holder_role", "visa_date")
REQ = {"company": "Denumirea firmei", "cui": "CUI", "contact_name": "Persoană de contact", "contact_phone": "Telefon", "contact_email": "E-mail",
       "attest_no": "Numărul atestatului ANRE", "holder_name": "Titularul atestatului (firma)"}
now = lambda: datetime.now(timezone.utc).isoformat()


def owner_only(user):
    if user.get("role") != "admin" or user["uid"] != user["id"]:
        raise HTTPException(status_code=403, detail="Doar titularul contului firmei poate gestiona atestatul")


def platform_only(user):
    if user["uid"] != S.PLATFORM["id"]:
        raise HTTPException(status_code=403, detail="Doar administratorul platformei are acces la această secțiune")


def pub(o):
    return {k: v for k, v in (o or {}).items() if k not in ("_id", "doc_path")} | {"has_document": bool((o or {}).get("doc_path"))}


@api.get("/operator")
async def op_get(user: dict = Depends(get_current_user)):
    o = await db.operators.find_one({"tenant": user["id"]}, {"_id": 0})
    return {"operator": pub(o) if o else None, "types": TYPES, "attest": user["plan"].get("attest")}


@api.put("/operator")
async def op_save(request: Request, user: dict = Depends(get_current_user)):
    owner_only(user)
    b = await request.json()
    d = {k: b.get(k) for k in FIELDS}
    d["attest_types"] = [t for t in (d.get("attest_types") or []) if t in TYPES]
    miss = [l for k, l in REQ.items() if not str(d.get(k) or "").strip()]
    if not d["attest_types"] and not (d.get("attest_other") or "").strip():
        miss.append("Tipul atestatului")
    if miss:
        raise HTTPException(status_code=400, detail="Completați: " + ", ".join(miss))
    cur = await db.operators.find_one({"tenant": user["id"]}, {"_id": 0}) or {}
    changed = any(cur.get(k) != d.get(k) for k in ("cui", "attest_no", "attest_types", "attest_other", "attest_valid", "holder_name"))
    st = "în așteptare" if changed or not cur.get("status") else cur["status"]
    if d.get("visa_date"):
        try:
            d["visa_date"] = date.fromisoformat(str(d["visa_date"]).strip()).isoformat()
        except ValueError:
            raise HTTPException(status_code=400, detail="Data vizării periodice invalidă (AAAA-LL-ZZ)")
    if cur.get("visa_date") != d.get("visa_date"):
        d["visa_confirmed"] = False
    await db.operators.update_one({"tenant": user["id"]}, {"$set": {**d, "tenant": user["id"], "email": user.get("email"), "status": st, "updated_at": now()},
                                                          "$push": {"history": {"at": now(), "by": user.get("email"), "event": "date atestat trimise" + (" — reverificare necesară" if changed and cur else "")}}}, upsert=True)
    await S.audit(user, "atestat ANRE trimis", "operators", user["id"], [st])
    return await op_get(user)


@api.post("/operator/document")
async def op_doc(file: UploadFile = File(...), user: dict = Depends(get_current_user)):
    owner_only(user)
    data = await file.read()
    ct = file.content_type or ""
    magic = {"application/pdf": b"%PDF", "image/jpeg": b"\xff\xd8\xff", "image/png": b"\x89PNG\r\n\x1a\n"}
    if ct not in magic or not data.startswith(magic[ct]) or len(data) > 10 * 1024 * 1024 or (ct != "application/pdf" and len(data) < 2000):
        raise HTTPException(status_code=400, detail="Încărcați atestatul ca PDF, JPG sau PNG, maxim 10 MB")
    st = S.put_object(f"{S.APP_NAME}/operators/{user['id']}/{datetime.now(timezone.utc).strftime('%Y%m%d%H%M%S')}_{(file.filename or 'atestat')[-60:]}", data, ct)
    await db.operators.update_one({"tenant": user["id"]}, {"$set": {"tenant": user["id"], "doc_path": st["path"], "doc_name": file.filename, "doc_type": ct, "doc_size": len(data),
                                                                    "doc_uploaded_at": now(), "status": "în așteptare", "updated_at": now()},
                                                          "$push": {"history": {"at": now(), "by": user.get("email"), "event": f"document încărcat: {file.filename}"}}}, upsert=True)
    await S.audit(user, "document atestat încărcat", "operators", user["id"], [file.filename or ""])
    return await op_get(user)


async def send_doc(tenant: str):
    o = await db.operators.find_one({"tenant": tenant}, {"_id": 0})
    if not o or not o.get("doc_path"):
        raise HTTPException(status_code=404, detail="Document inexistent")
    data, _ = S.get_object(o["doc_path"])
    return Response(content=data, media_type=o.get("doc_type") or "application/pdf", headers={"Content-Disposition": S.disp("inline", o.get("doc_name") or "atestat.pdf"), "Cache-Control": "no-store"})


@api.get("/operator/document")
async def op_doc_get(authorization: str = Header(None), auth: str = Query(None)):
    user = await S.user_from_token(authorization, auth)
    return await send_doc(user["id"])


@api.get("/admin/operators")
async def adm_list(user: dict = Depends(get_current_user)):
    platform_only(user)
    return [pub(o) for o in await db.operators.find({}, {"_id": 0}).sort("updated_at", -1).to_list(2000)]


@api.get("/admin/operators/{tenant}/document")
async def adm_doc(tenant: str, authorization: str = Header(None), auth: str = Query(None)):
    user = await S.user_from_token(authorization, auth)
    platform_only(user)
    await S.audit(user, "vizualizare atestat", "operators", tenant, [])
    return await send_doc(tenant)


@api.post("/admin/operators/{tenant}/decision")
async def adm_decide(tenant: str, request: Request, user: dict = Depends(get_current_user)):
    platform_only(user)
    b = await request.json()
    st, notes = b.get("status"), str(b.get("notes") or "").strip()[:1000]
    if st not in STATUSES:
        raise HTTPException(status_code=400, detail="Status invalid")
    o = await db.operators.find_one({"tenant": tenant})
    if not o:
        raise HTTPException(status_code=404, detail="Operator inexistent")
    if st == "aprobat" and not o.get("doc_path"):
        raise HTTPException(status_code=400, detail="Nu se poate aproba fără documentul atestatului încărcat și verificat")
    if st in ("respins", "expirat") and not notes:
        raise HTTPException(status_code=400, detail="Scrieți motivul în observații")
    await db.operators.update_one({"tenant": tenant}, {"$set": {"status": st, "admin_notes": notes, "verified_at": now(), "verified_by": user.get("email"), "updated_at": now()},
                                                       "$push": {"history": {"at": now(), "by": user.get("email"), "event": f"decizie: {st}" + (f" — {notes}" if notes else "")}}})
    await S.audit(user, f"atestat {st}", "operators", tenant, [notes])
    return {"status": st}


@api.post("/admin/operators/{tenant}/visa")
async def adm_visa(tenant: str, request: Request, user: dict = Depends(get_current_user)):
    platform_only(user)
    b = await request.json()
    o = await db.operators.find_one({"tenant": tenant}, {"_id": 0})
    if not o:
        raise HTTPException(status_code=404, detail="Operator inexistent")
    vd = str(b.get("visa_date") or o.get("visa_date") or "").strip()
    try:
        vd = date.fromisoformat(vd).isoformat()
    except ValueError:
        raise HTTPException(status_code=400, detail="Completați data vizării periodice (AAAA-LL-ZZ)")
    ok = bool(b.get("confirmed", True))
    await db.operators.update_one({"tenant": tenant}, {"$set": {"visa_date": vd, "visa_confirmed": ok, "visa_confirmed_by": user.get("email"), "visa_confirmed_at": now()},
                                                       "$push": {"history": {"at": now(), "by": user.get("email"), "event": f"dată vizare {'confirmată' if ok else 'neconfirmată'}: {vd}"}}})
    await S.audit(user, "dată vizare ANRE", "operators", tenant, [vd, str(ok)])
    return pub(await db.operators.find_one({"tenant": tenant}, {"_id": 0}))


@api.get("/admin/operators/visas")
async def adm_visas(user: dict = Depends(get_current_user)):
    platform_only(user)
    t = date.today()
    out = []
    async for o in db.operators.find({"visa_date": {"$nin": [None, ""], "$lte": (t + timedelta(days=90)).isoformat()}}, {"_id": 0, "doc_path": 0, "history": 0}).sort("visa_date", 1):
        out.append({k: o.get(k) for k in ("tenant", "company", "cui", "attest_no", "contact_name", "contact_email", "contact_phone", "visa_date", "visa_confirmed", "visa_reminded_at", "last_contact_at", "last_call", "status")}
                   | {"days_left": (date.fromisoformat(o["visa_date"]) - t).days})
    return out


VISA_COLS = (("company", "Firmă"), ("cui", "CUI"), ("attest_no", "Nr. atestat"), ("visa_date", "Data vizării"), ("days_left", "Zile rămase"), ("visa_confirmed", "Dată confirmată"),
             ("contact_name", "Persoană contact"), ("contact_phone", "Telefon"), ("contact_email", "E-mail"), ("visa_reminded_at", "Memento trimis"), ("last_contact_at", "Ultimul e-mail"),
             ("call_at", "Ultimul apel"), ("call_outcome", "Rezultat apel"), ("callback_date", "Data revenirii"), ("call_note", "Notiță apel"))


@api.get("/admin/operators/visas.xlsx")
async def adm_visas_xlsx(authorization: str = Header(None), auth: str = Query(None)):
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    user = await S.user_from_token(authorization, auth)
    rows = await adm_visas(user)
    wb = Workbook()
    ws = wb.active
    ws.title = "Vizări ANRE"
    ws.append([h for _, h in VISA_COLS] + ["Observații apel"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0B1F3A")
    for r in rows:
        lc = r.get("last_call") or {}
        r = {**r, "call_at": lc.get("at"), "call_outcome": lc.get("outcome"), "callback_date": lc.get("callback_date"), "call_note": lc.get("note")}
        ws.append([("da" if r.get(k) else "nu") if k == "visa_confirmed" else (str(r.get(k) or "")[:10] if k in ("visa_reminded_at", "last_contact_at", "visa_date", "call_at") else r.get(k)) for k, _ in VISA_COLS] + [""])
        if r["days_left"] < 0:
            for c in ws[ws.max_row]:
                c.fill = PatternFill("solid", fgColor="FDECEC")
    for col, w in zip("ABCDEFGHIJKLMNOP", (34, 12, 14, 13, 11, 14, 24, 16, 28, 14, 14, 13, 20, 14, 40, 40)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return Response(content=out.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": S.disp("attachment", f"vizari_ANRE_{date.today().isoformat()}.xlsx"), "Cache-Control": "no-store"})


def visa_prompt(o, instructions):
    import ai_routes as AI
    left = (date.fromisoformat(o["visa_date"]) - date.today()).days
    info = {"firma": o.get("company"), "persoana_contact": o.get("contact_name"), "nr_atestat": o.get("attest_no"), "tipuri_atestat": o.get("attest_types"),
            "data_vizarii": date.fromisoformat(o["visa_date"]).strftime("%d.%m.%Y"), "zile_ramase": left, "situatie": "termen depășit" if left < 0 else "termen apropiat"}
    return ("Redactează un e-mail personalizat de la DNS HOSTESS S.R.L. (platforma D-ENERGY) către firma de mai jos, despre vizarea periodică a atestatului ANRE. "
            "Ton profesionist, cordial, concis (max 180 de cuvinte). Menționează data și situația termenului, ce acte se pregătesc de obicei (certificat constatator recent, personal autorizat ANRE, dotări) "
            "și oferă serviciul D-ENERGY de pregătire a dosarului de vizare (3.200 lei + TVA; tarifele ANRE se plătesc separat). Nu inventa alte prețuri sau termene legale exacte. "
            "Format obligatoriu: prima linie „Subiect: ...”, apoi o linie goală, apoi corpul e-mailului semnat „Echipa D-ENERGY — DNS HOSTESS S.R.L.”.\n"
            f"Instrucțiuni suplimentare admin: {str(instructions or '—')[:600]}\nDate firmă (JSON): {AI.json.dumps(info, ensure_ascii=False)}")


SYS_MAIL = "Ești redactor de e-mailuri comerciale pentru D-ENERGY. "


@api.post("/admin/operators/{tenant}/ai-email")
async def adm_ai_email(tenant: str, request: Request, user: dict = Depends(get_current_user)):
    import ai_routes as AI
    platform_only(user)
    o = await db.operators.find_one({"tenant": tenant}, {"_id": 0, "doc_path": 0, "history": 0})
    if not o or not o.get("visa_date"):
        raise HTTPException(status_code=404, detail="Operatorul nu are dată de vizare")
    b = await request.json()
    return AI.sse(AI.new_chat(SYS_MAIL + AI.RULES), visa_prompt(o, b.get("instructions")))


async def send_op_email(tenant, subject, body, user):
    o = await db.operators.find_one({"tenant": tenant}, {"_id": 0, "contact_email": 1, "email": 1, "company": 1})
    if not o:
        raise HTTPException(status_code=404, detail="Operator inexistent")
    subj, body = str(subject or "").strip()[:200], str(body or "").strip()[:8000]
    to = o.get("contact_email") or o.get("email")
    if not subj or not body or not to:
        raise HTTPException(status_code=400, detail="Completați subiectul și textul e-mailului")
    html = '<div style="font-family:Arial,sans-serif;color:#0b1f3a">' + "".join(f"<p>{hesc(p).replace(chr(10), '<br/>')}</p>" for p in body.split("\n\n") if p.strip()) + "</div>"
    await send_email(to=to, subject=subj, html=html, reply_to=user.get("email"))
    await db.operators.update_one({"tenant": tenant}, {"$set": {"last_contact_at": now()}, "$push": {"history": {"at": now(), "by": user.get("email"), "event": f"e-mail trimis la {to}: {subj}"}}})
    await S.audit(user, "e-mail către operator", "operators", tenant, [to, subj])
    return to


@api.post("/admin/operators/{tenant}/email")
async def adm_send_email(tenant: str, request: Request, user: dict = Depends(get_current_user)):
    platform_only(user)
    b = await request.json()
    return {"sent": True, "to": await send_op_email(tenant, b.get("subject"), b.get("body"), user)}


def split_draft(t):
    m = re.match(r"\s*Subiect:\s*(.*)\n?([\s\S]*)$", t, re.I)
    return (m.group(1).strip(), m.group(2).strip()) if m else ("", t.strip())


@api.post("/admin/operators/campaign/drafts")
async def adm_campaign_drafts(request: Request, user: dict = Depends(get_current_user)):
    import ai_routes as AI
    platform_only(user)
    b = await request.json()
    t = date.today()
    recent = (datetime.now(timezone.utc) - timedelta(days=7)).isoformat()
    targets, skipped = [], []
    async for o in db.operators.find({"visa_date": {"$nin": [None, ""], "$lte": (t + timedelta(days=30)).isoformat()}}, {"_id": 0, "doc_path": 0, "history": 0, "calls": 0}).sort("visa_date", 1):
        why = "fără e-mail de contact" if not (o.get("contact_email") or o.get("email")) else "contactată în ultimele 7 zile" if (o.get("last_contact_at") or "") >= recent else None
        (skipped if why else targets).append({**o, "_why": why})
    sem = asyncio.Semaphore(4)

    async def one(o):
        async with sem:
            base = {"tenant": o["tenant"], "company": o.get("company"), "cui": o.get("cui"), "to": o.get("contact_email") or o.get("email"),
                    "visa_date": o["visa_date"], "days_left": (date.fromisoformat(o["visa_date"]) - t).days}
            try:
                subj, body = split_draft(re.sub(r"^#{1,6}\s+", "", re.sub(r"\*\*(.+?)\*\*", r"\1", await AI.collect(AI.new_chat(SYS_MAIL + AI.RULES), visa_prompt(o, b.get("instructions")))), flags=re.M))
                return {**base, "subject": subj, "body": body}
            except Exception as e:
                S.logger.error(f"campaign draft failed {o['tenant']}: {e}")
                return {**base, "subject": "", "body": "", "error": "Ciorna nu a putut fi redactată; scrieți textul manual sau reîncercați."}
    items = await asyncio.gather(*[one(o) for o in targets[:40]])
    return {"items": items, "skipped": [{"company": o.get("company"), "cui": o.get("cui"), "reason": o["_why"]} for o in skipped]}


@api.post("/admin/operators/campaign/send")
async def adm_campaign_send(request: Request, user: dict = Depends(get_current_user)):
    platform_only(user)
    items = (await request.json()).get("items") or []
    if not items or len(items) > 40:
        raise HTTPException(status_code=400, detail="Selectați între 1 și 40 de e-mailuri")
    res = []
    for i in items:
        try:
            res.append({"tenant": i.get("tenant"), "sent": True, "to": await send_op_email(str(i.get("tenant")), i.get("subject"), i.get("body"), user)})
        except HTTPException as e:
            res.append({"tenant": i.get("tenant"), "sent": False, "error": e.detail})
    await S.audit(user, "campanie vizări ANRE", "operators", "campaign", [f"{sum(r['sent'] for r in res)}/{len(res)} trimise"])
    return {"results": res, "sent": sum(r["sent"] for r in res)}


CALL_OUTCOMES = ("interesat", "nu răspunde", "revine / sună înapoi", "a comandat", "refuză", "alt rezultat")


@api.post("/admin/operators/{tenant}/calls")
async def adm_call(tenant: str, request: Request, user: dict = Depends(get_current_user)):
    platform_only(user)
    b = await request.json()
    if b.get("outcome") not in CALL_OUTCOMES:
        raise HTTPException(status_code=400, detail="Alegeți rezultatul apelului")
    cb = str(b.get("callback_date") or "").strip()
    if cb:
        try:
            cb = date.fromisoformat(cb).isoformat()
        except ValueError:
            raise HTTPException(status_code=400, detail="Data revenirii invalidă")
    c = {"id": str(uuid.uuid4()), "at": now(), "by": user.get("email"), "outcome": b["outcome"], "note": str(b.get("note") or "").strip()[:1000], "callback_date": cb or None}
    r = await db.operators.update_one({"tenant": tenant}, {"$push": {"calls": c}, "$set": {"last_call": c}})
    if not r.matched_count:
        raise HTTPException(status_code=404, detail="Operator inexistent")
    await S.audit(user, "apel operator", "operators", tenant, [c["outcome"]])
    return c


@api.get("/admin/operators/{tenant}/calls")
async def adm_calls(tenant: str, user: dict = Depends(get_current_user)):
    platform_only(user)
    o = await db.operators.find_one({"tenant": tenant}, {"_id": 0, "calls": 1})
    if o is None:
        raise HTTPException(status_code=404, detail="Operator inexistent")
    return {"items": list(reversed(o.get("calls") or [])), "outcomes": CALL_OUTCOMES}


async def run_visa(run_id: str, base: str):
    t = date.today()
    lim = (t + timedelta(days=60)).isoformat()
    async for o in db.operators.find({"visa_confirmed": True, "visa_date": {"$gte": t.isoformat(), "$lte": lim}}, {"_id": 0}):
        key = f"visa:{o['tenant']}:{o['visa_date']}"
        if await db.reminders.find_one({"id": key}):
            continue
        to = o.get("contact_email") or o.get("email")
        if not to:
            continue
        vd = date.fromisoformat(o["visa_date"])
        try:
            await send_email(to=to, subject=f"Memento: vizarea periodică a atestatului ANRE — {vd.strftime('%d.%m.%Y')}",
                             html=f'<div style="font-family:Arial,sans-serif;color:#0b1f3a"><p>Bună ziua,</p><p>Atestatul ANRE nr. <b>{hesc(str(o.get("attest_no") or "—"))}</b> al firmei '
                                  f'<b>{hesc(str(o.get("company") or ""))}</b> are termenul de vizare periodică la <b>{vd.strftime("%d.%m.%Y")}</b> (peste {(vd - t).days} zile).</p>'
                                  "<p>Vă recomandăm să pregătiți din timp dosarul de vizare (certificat constatator recent, personal autorizat, dotări). "
                                  "Puteți comanda serviciul D-ENERGY de pregătire a dosarului de vizare (3.200 lei + TVA):</p>"
                                  f'<p><a href="{base}/dosare-anre?tip=vizare" style="display:inline-block;background:#0b1f3a;color:#ffffff;padding:12px 22px;border-radius:6px;text-decoration:none;font-weight:bold">Comandă dosarul de vizare</a></p>'
                                  "<p style=\"font-size:12px;color:#5b6b80\">Butonul deschide aplicația cu „Vizare periodică” deja aleasă; comanda se plasează doar după confirmarea dvs.</p><p>Echipa D-ENERGY</p></div>")
        except HTTPException:
            continue
        await db.reminders.insert_one({"id": key, "at": now(), "run_id": run_id})
        await db.operators.update_one({"tenant": o["tenant"]}, {"$set": {"visa_reminded_at": now()}, "$push": {"history": {"at": now(), "by": "cron", "event": f"memento vizare trimis la {to}"}}})


@api.post("/cron/anre-visa")
async def cron_visa(request: Request, background: BackgroundTasks):
    # Cron endpoints must ack 2xx immediately; enqueue/background the actual work.
    h = request.headers.get("authorization", "")
    if not h.startswith("Bearer ") or not hmac.compare_digest(h[7:], os.environ.get("WEBHOOK_CRON_SECRET", "")):
        raise HTTPException(status_code=401, detail="Unauthorized")
    run_id = request.headers.get("x-webhook-id") or str(uuid.uuid4())
    if await db.cron_runs.find_one({"id": run_id}):
        return {"status": "duplicate"}
    await db.cron_runs.insert_one({"id": run_id, "at": now(), "name": "anre-visa"})
    host = request.headers.get("x-forwarded-host") or request.headers.get("host") or request.url.netloc
    background.add_task(run_visa, run_id, f"https://{host.split(',')[0].strip()}")
    return {"status": "accepted"}


@api.get("/admin/attest-config")
async def adm_cfg(user: dict = Depends(get_current_user)):
    platform_only(user)
    c = await S.billing_cfg()
    return {"attest_block_existing": bool(c.get("attest_block_existing")), "attest_cutoff": c.get("attest_cutoff")}


@api.put("/admin/attest-config")
async def adm_cfg_set(request: Request, user: dict = Depends(get_current_user)):
    platform_only(user)
    v = bool((await request.json()).get("attest_block_existing"))
    await db.billing_config.update_one({"id": "main"}, {"$set": {"attest_block_existing": v}}, upsert=True)
    await S.audit(user, "regulă atestate conturi existente", "operators", "config", ["blocare imediată" if v else "păstrare acces până la verificare"])
    return await adm_cfg(user)


@S.app.on_event("startup")
async def _cutoff():
    await db.billing_config.update_one({"id": "main", "attest_cutoff": {"$exists": False}}, {"$set": {"attest_cutoff": now()}})
    await db.operators.create_index("tenant", unique=True)
    if not await db.billing_config.find_one({"id": "main", "tariffs_2026_06": True}):
        for code, price in (("start", 500), ("pro", 900), ("full", 1250)):
            await db.plans.update_one({"code": code}, {"$set": {"price": price}})
        await db.billing_config.update_one({"id": "main"}, {"$set": {"tariffs_2026_06": True, "payment_mode": "transfer", "default_months": 1}}, upsert=True)
