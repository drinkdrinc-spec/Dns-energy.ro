# Iteration 28: visa e-mail campaign (AI drafts + bulk send), operator call log, visas.xlsx call columns
# Module-level functions on purpose: xdist loadscope keeps them on one worker, in order (shared state).
import io
import uuid
from datetime import date, timedelta, datetime, timezone

import pytest
import requests
from openpyxl import load_workbook
from pymongo import MongoClient

BASE = [l.split("=", 1)[1].strip() for l in open("/app/frontend/.env") if l.startswith("REACT_APP_BACKEND_URL=")][0]
API = BASE.rstrip("/") + "/api"
env = dict(l.strip().split("=", 1) for l in open("/app/backend/.env") if "=" in l)
mdb = MongoClient(env["MONGO_URL"].strip('"'))[env["DB_NAME"].strip('"')]
RID = uuid.uuid4().hex[:6]
T = date.today()
NOW = datetime.now(timezone.utc)
# key: (days to visa, contact_email, last_contact_at)
CASES = {
    "past": (-5, "delivered@resend.dev", None),
    "d30": (30, "delivered@resend.dev", None),
    "old": (15, "delivered@resend.dev", (NOW - timedelta(days=10)).isoformat()),
    "d31": (31, "delivered@resend.dev", None),
    "recent": (10, "delivered@resend.dev", (NOW - timedelta(days=2)).isoformat()),
    "nomail": (12, None, None),
}
TEN = {k: f"TEST_iter28_{k}_{RID}" for k in CASES}
CUI = {k: f"RO28{k}{RID}" for k in CASES}
ST = {}


def login(email, pw):
    r = requests.post(f"{API}/auth/login", json={"email": email, "password": pw})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}, r.json()["access_token"]


@pytest.fixture(scope="module")
def AH():
    return login("drinkdrinc@gmail.com", "DEnergy2026!")


@pytest.fixture(scope="module")
def H(AH):
    return AH[0]


@pytest.fixture(scope="module")
def U():
    email = f"sub_iter28_{RID}@example.com"
    bill = {"client_type": "pf", "full_name": "TEST I28", "address": "Str. 1", "locality": "Vaslui", "county": "Vaslui", "country": "România", "phone": "0700000000"}
    r = requests.post(f"{API}/auth/register", json={"email": email, "password": "Test2026!x", "name": "TEST", "billing": bill})
    assert r.status_code == 200, r.text
    uid = mdb.users.find_one({"email": email})["id"]
    ST["uid"] = uid
    mdb.users.update_one({"id": uid}, {"$set": {"status": "active"}})
    return login(email, "Test2026!x")[0]


@pytest.fixture(scope="module", autouse=True)
def seed():
    for k, (d, em, lc) in CASES.items():
        o = {"tenant": TEN[k], "company": f"TEST {k} 28 SRL", "cui": CUI[k], "attest_no": f"28/{k}", "contact_name": "TEST Ion Pop",
             "contact_phone": "0700", "status": "aprobat", "attest_types": ["A1"], "visa_date": (T + timedelta(days=d)).isoformat(),
             "visa_confirmed": True, "updated_at": T.isoformat()}
        if em:
            o["contact_email"] = em
        if lc:
            o["last_contact_at"] = lc
        mdb.operators.insert_one(o)
    yield
    mdb.operators.delete_many({"tenant": {"$regex": f"^TEST_iter28_.*_{RID}$"}})
    uid = ST.get("uid")
    if uid:
        for c in ("subscriptions", "invoices", "billing_profiles", "billing_history", "operators"):
            mdb[c].delete_many({"tenant": uid})
        mdb.users.delete_many({"id": uid})


# ---- campaign drafts
def test_drafts_targeting_and_content(H):
    r = requests.post(f"{API}/admin/operators/campaign/drafts", json={"instructions": "Menționează că putem începe imediat."}, headers=H, timeout=300)
    assert r.status_code == 200, r.text
    d = r.json()
    mine = {i["tenant"]: i for i in d["items"] if i["tenant"] in TEN.values()}
    assert set(mine) == {TEN["past"], TEN["d30"], TEN["old"]}, list(mine)
    for k in ("past", "d30", "old"):
        it = mine[TEN[k]]
        assert it["cui"] == CUI[k] and it["to"] == "delivered@resend.dev"
        assert it["days_left"] == CASES[k][0]
        assert not it.get("error"), it
        assert len(it["subject"].strip()) > 5 and "Subiect" not in it["subject"]
        assert len(it["body"]) > 150
        assert not it["body"].lower().startswith("subiect")
        low = (it["subject"] + it["body"]).lower()
        assert "viz" in low and ("atestat" in low or "anre" in low), it
        assert "**" not in it["body"]
    ST["drafts"] = mine
    sk = {s["cui"]: s["reason"] for s in d["skipped"]}
    assert sk.get(CUI["recent"]) == "contactată în ultimele 7 zile"
    assert sk.get(CUI["nomail"]) == "fără e-mail de contact"
    assert CUI["d31"] not in sk and CUI["d31"] not in [i["cui"] for i in d["items"]]


def test_drafts_403_subscriber(U):
    r = requests.post(f"{API}/admin/operators/campaign/drafts", json={}, headers=U)
    assert r.status_code == 403


def test_drafts_401_no_auth():
    assert requests.post(f"{API}/admin/operators/campaign/drafts", json={}).status_code in (401, 403)


# ---- campaign send
def test_send_empty_and_too_many(H):
    r = requests.post(f"{API}/admin/operators/campaign/send", json={"items": []}, headers=H)
    assert r.status_code == 400 and "40" in r.json()["detail"]
    r = requests.post(f"{API}/admin/operators/campaign/send", json={}, headers=H)
    assert r.status_code == 400
    r = requests.post(f"{API}/admin/operators/campaign/send", json={"items": [{"tenant": "x", "subject": "a", "body": "b"}] * 41}, headers=H)
    assert r.status_code == 400


def test_send_403_subscriber(U):
    r = requests.post(f"{API}/admin/operators/campaign/send", json={"items": [{"tenant": TEN["past"], "subject": "a", "body": "b"}]}, headers=U)
    assert r.status_code == 403


def test_send_partial_failures(H):
    dr = ST["drafts"]
    items = [
        {"tenant": TEN["past"], "subject": dr[TEN["past"]]["subject"], "body": dr[TEN["past"]]["body"]},
        {"tenant": TEN["d30"], "subject": "", "body": dr[TEN["d30"]]["body"]},  # empty subject -> fail
        {"tenant": "TEST_iter28_unknown", "subject": "x", "body": "y"},  # unknown -> fail
        {"tenant": TEN["old"], "subject": "TEST_ " + dr[TEN["old"]]["subject"], "body": dr[TEN["old"]]["body"]},
    ]
    r = requests.post(f"{API}/admin/operators/campaign/send", json={"items": items}, headers=H, timeout=120)
    assert r.status_code == 200, r.text
    d = r.json()
    res = {x["tenant"]: x for x in d["results"]}
    assert d["sent"] == 2 and len(d["results"]) == 4
    assert res[TEN["past"]]["sent"] is True and res[TEN["past"]]["to"] == "delivered@resend.dev"
    assert res[TEN["old"]]["sent"] is True
    assert res[TEN["d30"]]["sent"] is False and "subiect" in res[TEN["d30"]]["error"].lower()
    assert res["TEST_iter28_unknown"]["sent"] is False and "inexistent" in res["TEST_iter28_unknown"]["error"].lower()
    for k in ("past", "old"):
        o = mdb.operators.find_one({"tenant": TEN[k]})
        assert o["last_contact_at"] >= (NOW - timedelta(minutes=5)).isoformat()
        assert any("e-mail trimis" in h["event"] for h in o.get("history", []))
    assert "last_contact_at" not in mdb.operators.find_one({"tenant": TEN["d30"]})


def test_visas_shows_last_contact(H):
    v = {x["tenant"]: x for x in requests.get(f"{API}/admin/operators/visas", headers=H).json()}
    assert v[TEN["past"]]["last_contact_at"] and v[TEN["d30"]].get("last_contact_at") is None


# ---- calls
def test_call_validation(H):
    t = TEN["d30"]
    r = requests.post(f"{API}/admin/operators/{t}/calls", json={"outcome": "nimic"}, headers=H)
    assert r.status_code == 400
    r = requests.post(f"{API}/admin/operators/{t}/calls", json={}, headers=H)
    assert r.status_code == 400
    r = requests.post(f"{API}/admin/operators/{t}/calls", json={"outcome": "interesat", "callback_date": "31-12-2026"}, headers=H)
    assert r.status_code == 400 and "revenirii" in r.json()["detail"]
    r = requests.post(f"{API}/admin/operators/TEST_iter28_nope/calls", json={"outcome": "interesat"}, headers=H)
    assert r.status_code == 404
    assert requests.get(f"{API}/admin/operators/TEST_iter28_nope/calls", headers=H).status_code == 404


def test_call_403_subscriber(U):
    t = TEN["d30"]
    assert requests.post(f"{API}/admin/operators/{t}/calls", json={"outcome": "interesat"}, headers=U).status_code == 403
    assert requests.get(f"{API}/admin/operators/{t}/calls", headers=U).status_code == 403


def test_call_log_history_and_last_call(H):
    t = TEN["d30"]
    r0 = requests.get(f"{API}/admin/operators/{t}/calls", headers=H).json()
    assert r0["items"] == [] and list(r0["outcomes"]) == ["interesat", "nu răspunde", "revine / sună înapoi", "a comandat", "refuză", "alt rezultat"]
    a = requests.post(f"{API}/admin/operators/{t}/calls", json={"outcome": "nu răspunde"}, headers=H)
    assert a.status_code == 200 and a.json()["callback_date"] is None and a.json()["by"] == "drinkdrinc@gmail.com"
    cb = (T - timedelta(days=1)).isoformat()
    b = requests.post(f"{API}/admin/operators/{t}/calls", json={"outcome": "revine / sună înapoi", "note": "TEST_ sună joi", "callback_date": cb}, headers=H)
    assert b.status_code == 200, b.text
    h = requests.get(f"{API}/admin/operators/{t}/calls", headers=H).json()["items"]
    assert [c["outcome"] for c in h] == ["revine / sună înapoi", "nu răspunde"]
    assert h[0]["id"] == b.json()["id"]
    v = {x["tenant"]: x for x in requests.get(f"{API}/admin/operators/visas", headers=H).json()}
    lc = v[t]["last_call"]
    assert lc["outcome"] == "revine / sună înapoi" and lc["callback_date"] == cb and lc["note"] == "TEST_ sună joi"
    assert "calls" not in v[t]


# ---- excel
def test_xlsx_call_columns(AH):
    r = requests.get(f"{API}/admin/operators/visas.xlsx", params={"auth": AH[1]})
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    hdr = [c.value for c in ws[1]]
    for h in ("Ultimul apel", "Rezultat apel", "Data revenirii", "Notiță apel", "Ultimul e-mail"):
        assert h in hdr, hdr
    rows = {row[hdr.index("CUI")]: row for row in ws.iter_rows(min_row=2, values_only=True)}
    row = rows[CUI["d30"]]
    assert row[hdr.index("Rezultat apel")] == "revine / sună înapoi"
    assert row[hdr.index("Data revenirii")] == (T - timedelta(days=1)).isoformat()
    assert row[hdr.index("Notiță apel")] == "TEST_ sună joi"
    assert row[hdr.index("Ultimul apel")] == NOW.date().isoformat()
    assert rows[CUI["past"]][hdr.index("Ultimul e-mail")] == datetime.now(timezone.utc).date().isoformat()
    ST["hdr"] = hdr


def test_xlsx_403_subscriber(U):
    tok = U["Authorization"].split(" ")[1]
    assert requests.get(f"{API}/admin/operators/visas.xlsx", params={"auth": tok}).status_code == 403


# ---- second drafts call skips freshly contacted
def test_drafts_after_send_skips_contacted(H):
    r = requests.post(f"{API}/admin/operators/campaign/drafts", json={}, headers=H, timeout=300)
    assert r.status_code == 200, r.text
    d = r.json()
    tenants = [i["tenant"] for i in d["items"]]
    assert TEN["past"] not in tenants and TEN["old"] not in tenants
    assert TEN["d30"] in tenants
    sk = {s["cui"]: s["reason"] for s in d["skipped"]}
    assert sk.get(CUI["past"]) == "contactată în ultimele 7 zile" and sk.get(CUI["old"]) == "contactată în ultimele 7 zile"
