import os as _dns_test_os
import pytest as _dns_test_pytest
if not (_dns_test_os.environ.get("DNS_ENERGY_TEST_EMAIL") and _dns_test_os.environ.get("DNS_ENERGY_TEST_PASSWORD") and _dns_test_os.environ.get("REACT_APP_BACKEND_URL") and _dns_test_os.environ.get("DNS_ENERGY_TEST_ALLOW_MUTATION") == "1"):
    _dns_test_pytest.skip("Set staging-only DNS_ENERGY_TEST_* credentials, REACT_APP_BACKEND_URL and DNS_ENERGY_TEST_ALLOW_MUTATION=1", allow_module_level=True)

# Iteration 21: DI-U remote signing (one link per doc), signed PDF propagation, Word export, DI-U AI precheck, sign-message module=diu
import os, io, hashlib, base64, random
from datetime import datetime, timezone, timedelta
import pytest, requests
from pymongo import MongoClient
from PIL import Image
from pypdf import PdfReader

BASE = os.environ["REACT_APP_BACKEND_URL"].rstrip("/") + "/api"
DID = "70b56325-ab31-42e3-84f2-c115d6c6ba78"
env = dict(l.strip().split("=", 1) for l in open("/app/backend/.env") if "=" in l)
mdb = MongoClient(env["MONGO_URL"].strip('"'))[env["DB_NAME"].strip('"')]
ALL = [k for k, _ in __import__("diu_doc").DOCS]
ST = {}


@pytest.fixture(scope="module")
def H():
    r = requests.post(f"{BASE}/auth/login", json={"email": _dns_test_os.environ["DNS_ENERGY_TEST_EMAIL"], "password": _dns_test_os.environ["DNS_ENERGY_TEST_PASSWORD"]})
    assert r.status_code == 200
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


@pytest.fixture(scope="module", autouse=True)
def cleanup():
    before_docs = mdb.documents.count_documents({"dossier_id": DID})
    yield
    ids = ST.get("req_ids", [])
    mdb.sign_requests.delete_many({"id": {"$in": ids}})
    mdb.diu.update_one({"id": DID}, {"$unset": {"signed": "", "ai_precheck": ""}})
    mdb.documents.delete_many({"dossier_id": DID, "kind": "Semnat electronic"})
    cid = mdb.diu.find_one({"id": DID})["client_id"]
    mdb.clients.update_one({"id": cid}, {"$pull": {"docs.semnate": {"original_filename": {"$regex": "^DI-U"}}}})


def test_docs_list_has_11_and_docx_flags(H):
    r = requests.get(f"{BASE}/diu/{DID}/docs", headers=H)
    assert r.status_code == 200
    pos = r.json()["positions"]
    keys = [p["key"] for p in pos]
    print("keys", keys)
    assert len(pos) == 11
    ST["keys"] = keys
    for p in pos:
        assert p["docx"] == (p["key"] != "schema_tg")


def test_sign_create_all_11(H):
    keys = ST["keys"]
    r = requests.post(f"{BASE}/diu/{DID}/sign", headers=H, json={"pos_keys": keys, "signer": {"name": "TEST Client", "email": "test_i21@example.com", "phone": "0700"}, "expires_days": 7})
    assert r.status_code == 200, r.text
    out = r.json()
    assert len(out) == 11
    assert len({x["token"] for x in out}) == 11
    for x, k in zip(out, keys):
        assert x["pos_key"] == k
        assert x["doc_title"].startswith("DI-U — ")
    ST["req_ids"] = [x["id"] for x in out]
    ST["tok"] = {x["pos_key"]: x["token"] for x in out}
    for x in out:
        d = mdb.sign_requests.find_one({"id": x["id"]})
        assert d["module"] == "diu" and d["dossier_id"] == DID
    assert len({mdb.sign_requests.find_one({"id": i})["group_id"] for i in ST["req_ids"]}) == 11


def test_sign_invalid_key_and_email(H):
    r = requests.post(f"{BASE}/diu/{DID}/sign", headers=H, json={"pos_keys": ["nope"], "signer": {"name": "X", "email": "a@b.ro"}})
    assert r.status_code == 400
    r = requests.post(f"{BASE}/diu/{DID}/sign", headers=H, json={"pos_keys": ["garantie"], "signer": {"name": "X", "email": "bad"}})
    assert r.status_code == 400
    r = requests.post(f"{BASE}/diu/{DID}/sign", headers=H, json={"pos_keys": [], "signer": {"name": "X", "email": "a@b.ro"}})
    assert r.status_code == 400


def test_sign_list_module_diu(H):
    r = requests.get(f"{BASE}/sign/requests", headers=H, params={"dossier_id": DID, "module": "diu"})
    assert r.status_code == 200
    ids = {x["id"] for x in r.json()}
    assert set(ST["req_ids"]) <= ids
    r2 = requests.get(f"{BASE}/sign/requests", headers=H, params={"dossier_id": DID, "module": "bad"})
    assert r2.status_code == 400


def _png():
    img = Image.new("RGB", (300, 120), "white"); px = img.load()
    for _ in range(4000):
        px[random.randrange(300), random.randrange(120)] = (0, 0, random.randrange(255))
    b = io.BytesIO(); img.save(b, "PNG")
    return "data:image/png;base64," + base64.b64encode(b.getvalue()).decode()


def test_public_sign_garantie(H):
    tok = ST["tok"]["garantie"]
    pub = requests.get(f"{BASE}/public/sign/{tok}")
    assert pub.status_code == 200, pub.text
    print("public view keys", list(pub.json().keys()))
    req = mdb.sign_requests.find_one({"token_hash": hashlib.sha256(tok.encode()).hexdigest()})
    mdb.sign_requests.update_one({"id": req["id"]}, {"$set": {"otp_hash": hashlib.sha256(f"{req['id']}:123456".encode()).hexdigest(),
                                                              "otp_exp": (datetime.now(timezone.utc) + timedelta(minutes=10)).isoformat(), "otp_attempts": 0}})
    pros_before = mdb.prosumator.find_one({"id": mdb.diu.find_one({"id": DID})["pros_id"]}, {"_id": 0, "opis": 1})
    cid = mdb.diu.find_one({"id": DID})["client_id"]
    n_sem = len(((mdb.clients.find_one({"id": cid}) or {}).get("docs") or {}).get("semnate") or [])
    r = requests.post(f"{BASE}/public/sign/{tok}/sign", json={"otp": "123456", "intent": True, "declared_name": "TEST Client", "signature": _png()})
    assert r.status_code == 200, r.text
    d = mdb.diu.find_one({"id": DID})
    assert d["signed"]["garantie"]["signed_at"]
    assert d["signed"]["garantie"]["storage_path"]
    assert mdb.prosumator.find_one({"id": d["pros_id"]}, {"_id": 0, "opis": 1}) == pros_before
    sem = mdb.clients.find_one({"id": cid})["docs"]["semnate"]
    assert len(sem) == n_sem + 1 and "garan" in sem[-1]["original_filename"].lower()
    doc = mdb.documents.find_one({"dossier_id": DID, "kind": "Semnat electronic"})
    assert doc and doc["doc_type"].startswith("DI-U — ")


def test_docs_status_semnat(H):
    pos = {p["key"]: p for p in requests.get(f"{BASE}/diu/{DID}/docs", headers=H).json()["positions"]}
    assert pos["garantie"]["status"] == "semnat" and pos["garantie"]["signed_at"]
    assert pos["opis"]["status"] != "semnat"


def test_pdf_signed_vs_orig(H):
    s = requests.get(f"{BASE}/diu/{DID}/pdf", headers=H, params={"doc": "garantie"})
    o = requests.get(f"{BASE}/diu/{DID}/pdf", headers=H, params={"doc": "garantie", "orig": 1})
    assert s.status_code == 200 and o.status_code == 200
    ps, po = PdfReader(io.BytesIO(s.content)), PdfReader(io.BytesIO(o.content))
    ts = " ".join(" ".join((p.extract_text() or "").split()) for p in ps.pages)
    print("signed pages", len(ps.pages), "orig pages", len(po.pages), s.headers.get("content-disposition"))
    assert len(ps.pages) > len(po.pages)
    assert "Semnat electronic" in ts
    a = requests.get(f"{BASE}/diu/{DID}/pdf", headers=H, params={"doc": "all"})
    assert a.status_code == 200
    pa = PdfReader(io.BytesIO(a.content))
    ta = " ".join(" ".join((p.extract_text() or "").split()) for p in pa.pages)
    assert "Semnat electronic" in ta


@pytest.mark.parametrize("k", [k for k in ALL if k != "schema_tg"])
def test_docx_each(H, k):
    from docx import Document
    r = requests.get(f"{BASE}/diu/{DID}/pdf", headers=H, params={"doc": k, "fmt": "docx"})
    assert r.status_code == 200, r.text[:200]
    assert "wordprocessingml" in r.headers["content-type"]
    d = Document(io.BytesIO(r.content))
    txt = "\n".join(p.text for p in d.paragraphs) + "\n".join(c.text for t in d.tables for row in t.rows for c in row.cells)
    assert len(txt.strip()) > 50
    titles = {p["key"]: p["title"] for p in requests.get(f"{BASE}/diu/{DID}/docs", headers=H).json()["positions"]}
    first = titles[k].split()[0].upper()
    assert first in txt.upper(), (k, titles[k], txt[:300])


def test_docx_rejected(H):
    for k in ("schema_tg", "all"):
        r = requests.get(f"{BASE}/diu/{DID}/pdf", headers=H, params={"doc": k, "fmt": "docx"})
        assert r.status_code == 400


def test_put_does_not_wipe_signed(H):
    # Frontend sends the whole form (stale copy) on save -> server-managed 'signed' must survive
    cur = requests.get(f"{BASE}/data/diu/{DID}", headers=H).json()
    stale = {k: v for k, v in cur.items() if k != "_id"}
    stale["signed"] = {}
    r = requests.put(f"{BASE}/data/diu/{DID}", headers=H, json=stale)
    assert r.status_code == 200
    after = mdb.diu.find_one({"id": DID})
    ok = bool((after.get("signed") or {}).get("garantie"))
    if not ok:
        mdb.diu.update_one({"id": DID}, {"$set": {"signed": cur.get("signed")}})
    assert ok, "PUT /data/diu overwrote server-managed diu.signed with client copy"


def test_ai_diu_precheck_and_ack(H):
    r = requests.post(f"{BASE}/ai/diu/{DID}/precheck", headers=H, timeout=180)
    assert r.status_code == 200, r.text
    pc = r.json()
    print("items", len(pc["items"]), pc["items"][:3])
    assert isinstance(pc["items"], list)
    for i in pc["items"]:
        assert i["text"] and i["severity"] in ("critic", "important", "minor")
    assert mdb.diu.find_one({"id": DID})["ai_precheck"]["at"] == pc["at"]
    a = requests.post(f"{BASE}/ai/diu/{DID}/precheck/ack", headers=H, json={"checked": [0, 999]})
    assert a.status_code == 200
    assert a.json()["checked"] == ([0] if pc["items"] else [])


def test_ai_sign_message_diu(H):
    links = [{"title": "DI-U — Garanție", "url": "https://x.example/semnare/AAA111"}, {"title": "DI-U — Opis", "url": "https://x.example/semnare/BBB222"}]
    r = requests.post(f"{BASE}/ai/sign-message", headers=H, json={"dossier_id": DID, "module": "diu", "channel": "whatsapp", "links": links}, stream=True, timeout=180)
    assert r.status_code == 200
    import json
    out, err = "", None
    for line in r.iter_lines(decode_unicode=True):
        if line and line.startswith("data: "):
            j = json.loads(line[6:])
            out += j.get("t", ""); err = j.get("error") or err
            if j.get("done"):
                break
    print(out[:400])
    assert not err
    assert "AAA111" in out and "BBB222" in out


def test_prosumer_precheck_regression(H):
    pid = mdb.diu.find_one({"id": DID})["pros_id"]
    r = requests.post(f"{BASE}/ai/dossier/{pid}/precheck", headers=H, timeout=180)
    assert r.status_code == 200, r.text
    assert isinstance(r.json()["items"], list)
    mdb.prosumator.update_one({"id": pid}, {"$unset": {"ai_precheck": ""}})


def test_pro_plan_diu_ai_402_and_sign(H):
    # PRO plan (no ai_gpt) -> DI-U precheck/ack 402; DI-U sign gated by sign_remote
    import uuid
    PF = {"client_type": "pf", "full_name": "TEST Pro", "address": "Str. X 1", "locality": "Iasi", "county": "Iasi", "country": "Romania", "phone": "0700000000"}
    email = f"sub_i21_{uuid.uuid4().hex[:6]}@example.com"
    assert requests.post(f"{BASE}/auth/register", json={"email": email, "password": "Test2026!x", "name": "TEST pro", "billing": PF}).status_code == 200
    uid = mdb.users.find_one({"email": email})["id"]
    try:
        requests.put(f"{BASE}/admin/subscribers/{uid}", headers=H, json={"status": "active"})
        requests.post(f"{BASE}/admin/subscribers/{uid}/plan", headers=H, json={"plan": "pro", "valid_until": "2026-12-31", "status": "activ", "justification": "TEST"})
        t = requests.post(f"{BASE}/auth/login", json={"email": email, "password": "Test2026!x"}).json()["access_token"]
        h = {"Authorization": f"Bearer {t}"}
        did = requests.post(f"{BASE}/data/diu", headers=h, json={"number": "TEST_PRO_I21"}).json()["id"]
        r = requests.post(f"{BASE}/ai/diu/{did}/precheck", headers=h)
        assert r.status_code == 402, r.text
        r = requests.post(f"{BASE}/ai/diu/{did}/precheck/ack", headers=h, json={"checked": []})
        assert r.status_code == 402
        r = requests.post(f"{BASE}/diu/{did}/sign", headers=h, json={"pos_keys": ["garantie"], "signer": {"name": "X", "email": "a@b.ro"}})
        print("PRO diu sign status", r.status_code, r.text[:120])
        r = requests.get(f"{BASE}/diu/{did}/pdf", headers=h, params={"doc": "memoriu", "fmt": "docx"})
        assert r.status_code == 200
    finally:
        mdb.subscriptions.delete_many({"tenant": uid}); mdb.users.delete_many({"id": uid}); mdb.diu.delete_many({"owner": uid}); mdb.sign_requests.delete_many({"owner": uid})
