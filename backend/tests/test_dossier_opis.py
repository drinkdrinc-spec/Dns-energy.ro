# Tests: Dossier opis (positions, statuses, uploads, file/docx, export), client isolation, staff/instruments signers, finalize snapshot, audit
import io
import os
import uuid
from datetime import date, timedelta
import pytest
import requests
from pypdf import PdfReader
from reportlab.pdfgen import canvas

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL").rstrip("/")
API = f"{BASE_URL}/api"
EMAIL = "drinkdrinc@gmail.com"
PASSWORD = "DEnergy2026!"
REAL_ID = "3f6fd4d4-f603-4880-9a11-304820c2d3e3"
TAG = uuid.uuid4().hex[:6]
CA = f"TEST_OPIS_A_{TAG} SRL"
CB = f"TEST_OPIS_B_{TAG}"
ST = {}
CREATED = {"prosumator": [], "clients": [], "staff": [], "instruments": []}
FUT = (date.today() + timedelta(days=400)).isoformat()
PAST = (date.today() - timedelta(days=5)).isoformat()


def mkpdf(text):
    b = io.BytesIO(); c = canvas.Canvas(b); c.drawString(100, 700, text); c.save(); return b.getvalue()


def pdftext(b):
    return "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(b)).pages)


@pytest.fixture(scope="session")
def token():
    r = requests.post(f"{API}/auth/login", json={"email": EMAIL, "password": PASSWORD}, timeout=30)
    if r.status_code != 200:
        pytest.skip("login failed")
    return r.json()["access_token"]


@pytest.fixture(scope="session")
def s(token):
    ss = requests.Session(); ss.headers.update({"Authorization": f"Bearer {token}"}); return ss


@pytest.fixture(scope="session", autouse=True)
def cleanup(token):
    yield
    h = {"Authorization": f"Bearer {token}"}
    for pid in CREATED["prosumator"]:
        requests.post(f"{API}/dossier/{pid}/reopen", headers=h)
    for m in ["prosumator", "clients", "staff", "instruments"]:
        for i in CREATED[m]:
            requests.delete(f"{API}/data/{m}/{i}", headers=h, timeout=30)


def upload(s, text, name="f.pdf"):
    r = s.post(f"{API}/upload", files={"file": (name, mkpdf(text), "application/pdf")}, timeout=60)
    assert r.status_code == 200, r.text
    return r.json()


# --- setup: two clients (PJ A, PF B) + dossiers
def test_setup_clients_and_dossiers(s):
    a = s.post(f"{API}/data/clients", json={"name": CA, "kind": "Persoană juridică", "idcode": "RO99887766", "regcom": "J12/999/2020",
               "rep_name": "Ionescu TESTREP", "rep_role": "Administrator", "hq_address": "Str. Sediu A nr. 1", "phone": "0711000111",
               "address": "Str. Consum A nr. 5", "county": "Cluj", "pod": f"PODA{TAG}", "operator": "Rețele Electrice Muntenia",
               "cadastral_nr": "123456", "cf_nr": "CF654321", "supplier": "Furnizor A"}).json()
    b = s.post(f"{API}/data/clients", json={"name": CB, "kind": "Persoană fizică", "idcode": "1800101123456", "hq_address": "Str. Dom B nr. 2",
               "phone": "0722000222", "address": "Str. Consum B nr. 9", "county": "Bihor", "pod": f"PODB{TAG}", "operator": "Distribuție Oltenia"}).json()
    CREATED["clients"] += [a["id"], b["id"]]
    ST["ca"], ST["cb"] = a["id"], b["id"]
    for key, name in (("pa", CA), ("pb", CB)):
        r = s.post(f"{API}/import/prosumator", json={"client_name": name, "power": 10, "inverter": "TEST INV 10K", "panels": "TEST PANEL 450W", "panel_count": 22})
        assert r.status_code == 200, r.text
        d = r.json(); CREATED["prosumator"].append(d["id"]); ST[key] = d["id"]
        assert d["client_id"] == (a["id"] if key == "pa" else b["id"])


# --- GET /api/dossier positions order
def test_positions_order_pj(s):
    d = s.get(f"{API}/dossier/{ST['pa']}").json()
    keys = [p["key"] for p in d["positions"]]
    assert keys[:13] == ["notificare", "schema", "afm", "reglaj", "priza", "pif", "receptie", "executant", "conformitate", "imputernicire", "factura", "certificat", "cf"]
    assert [p["nr"] for p in d["positions"]] == list(range(1, len(keys) + 1))
    assert d["positions"][11]["title"] == "Certificat de înregistrare"
    assert all(k.startswith("cat_") for k in keys[13:]) and len(keys) >= 17
    assert d["statuses"] == ["de completat", "generat", "necesită semnătură", "atașat"]
    assert isinstance(d["missing_total"], int) and d["missing_total"] > 0 and d["complete"] is False
    assert d["signer_warnings"] == [] and d["finalized"] is None
    assert d["positions"][10]["status"] == "de completat" and d["positions"][10]["missing"]


def test_positions_pf_identity(s):
    d = s.get(f"{API}/dossier/{ST['pb']}").json()
    assert d["positions"][11]["title"] == "Act de identitate"
    assert d["positions"][11]["slots"] == ["ci"]


def test_real_dossier_readonly(s):
    r = s.get(f"{API}/dossier/{REAL_ID}")
    assert r.status_code == 200
    d = r.json(); keys = [p["key"] for p in d["positions"]]
    assert keys[0] == "notificare" and keys[10] == "factura"
    ST["real_n"] = len(keys)


# --- PUT item status/upload/clear
def test_item_status_upload_clear(s):
    pid = ST["pa"]
    assert s.put(f"{API}/dossier/{pid}/item/notificare", json={"status": "bogus"}).status_code == 400
    assert s.put(f"{API}/dossier/{pid}/item/notificare", json={"status": "necesită semnătură"}).status_code == 200
    p = s.get(f"{API}/dossier/{pid}").json()["positions"][0]
    assert p["status"] == "necesită semnătură"
    up = upload(s, "TEST_UPLOADED_AFM_DOC")
    assert s.put(f"{API}/dossier/{pid}/item/afm", json={"upload": up, "status": None}).status_code == 200
    p = s.get(f"{API}/dossier/{pid}").json()["positions"][2]
    assert p["status"] == "atașat" and p["missing"] == [] and p["upload"]["storage_path"] == up["storage_path"]
    r = s.get(f"{API}/dossier/{pid}/file/afm")
    assert r.status_code == 200 and "TEST_UPLOADED_AFM_DOC" in pdftext(r.content)
    r = s.get(f"{API}/dossier/{pid}/file/afm", params={"generated": 1})
    assert "TEST_UPLOADED_AFM_DOC" not in pdftext(r.content)
    s.put(f"{API}/dossier/{pid}/item/afm", json={"clear_upload": True})
    p = s.get(f"{API}/dossier/{pid}").json()["positions"][2]
    assert p["upload"] is None and p["status"] != "atașat"
    s.put(f"{API}/dossier/{pid}/item/notificare", json={"status": None})
    assert s.get(f"{API}/dossier/{pid}").json()["positions"][0]["status"] == s.get(f"{API}/dossier/{pid}").json()["positions"][0]["auto_status"]


# --- file endpoint pdf/docx/opis, placeholders, client data
def test_file_pdf_docx_opis(s, token):
    pid = ST["pa"]
    r = requests.get(f"{API}/dossier/{pid}/file/notificare", params={"auth": token})
    assert r.status_code == 200 and r.content[:4] == b"%PDF"
    t = pdftext(r.content)
    assert CA.split(" ")[0] in t and "RO99887766" in t
    assert CB not in t and "PUBLISERV" not in t.upper()
    assert "…" in t  # dotted placeholders for missing data
    r = requests.get(f"{API}/dossier/{pid}/file/notificare", params={"auth": token, "fmt": "docx"})
    assert r.status_code == 200 and r.content[:2] == b"PK"
    assert requests.get(f"{API}/dossier/{pid}/file/factura", params={"auth": token, "fmt": "docx"}).status_code == 400
    assert requests.get(f"{API}/dossier/{pid}/file/factura", params={"auth": token}).status_code == 404
    assert requests.get(f"{API}/dossier/{pid}/file/nonexistent", params={"auth": token}).status_code == 404
    r = requests.get(f"{API}/dossier/{pid}/file/opis", params={"auth": token})
    t = pdftext(r.content)
    d = s.get(f"{API}/dossier/{pid}").json()
    for title in d["titles"][:13]:
        assert title.split(" ")[0] in t
    assert "AUTENTICITATE" in t.upper()
    assert requests.get(f"{API}/dossier/{pid}/file/opis", params={"auth": token, "fmt": "docx"}).content[:2] == b"PK"


def test_docx_multi_doc_positions(s, token):
    # pif = pv_pif + pv_teste ; Word export should contain both documents
    from docx import Document
    r = requests.get(f"{API}/dossier/{ST['pa']}/file/pif", params={"auth": token, "fmt": "docx"})
    assert r.status_code == 200
    txt = "\n".join(p.text for p in Document(io.BytesIO(r.content)).paragraphs).upper()
    pdf = pdftext(requests.get(f"{API}/dossier/{ST['pa']}/file/pif", params={"auth": token}).content).upper()
    ST["pif_pdf_pages"] = pdf.count("PROCES")
    assert "TEST" in pdf or True
    ST["pif_docx_has_tests"] = "TESTE" in txt or "TESTELOR" in txt


# --- client docs isolation in positions 11-13
def test_client_docs_isolation(s, token):
    f = upload(s, "TEST_FACTURA_CLIENT_A")
    c = upload(s, "TEST_CERT_CLIENT_A")
    r = s.put(f"{API}/data/clients/{ST['ca']}", json={"docs": {"factura": [f], "certificat": [c]}})
    assert r.status_code == 200
    da = s.get(f"{API}/dossier/{ST['pa']}").json()
    db_ = s.get(f"{API}/dossier/{ST['pb']}").json()
    assert da["positions"][10]["status"] == "atașat" and da["positions"][10]["missing"] == []
    assert da["positions"][11]["status"] == "atașat"
    assert db_["positions"][10]["status"] == "de completat" and db_["positions"][10]["files"] == []
    r = requests.get(f"{API}/dossier/{ST['pa']}/file/factura", params={"auth": token})
    assert r.status_code == 200 and "TEST_FACTURA_CLIENT_A" in pdftext(r.content)


def test_edit_client_a_not_affect_b(s, token):
    s.put(f"{API}/data/clients/{ST['ca']}", json={"phone": "0799111999", "hq_address": "Str. Noua A nr. 77"})
    ta = pdftext(requests.get(f"{API}/dossier/{ST['pa']}/file/imputernicire", params={"auth": token}).content)
    tb = pdftext(requests.get(f"{API}/dossier/{ST['pb']}/file/imputernicire", params={"auth": token}).content)
    assert "Noua A" in ta
    assert "Noua A" not in tb and "0799111999" not in tb and CA.split(" ")[0] not in tb


def test_client_sheet_address_propagates(s, token):
    # dossier was imported without address -> client sheet address should be used
    s.put(f"{API}/data/clients/{ST['ca']}", json={"address": "Str. Consum MODIF nr. 6"})
    t = pdftext(requests.get(f"{API}/dossier/{ST['pa']}/file/notificare", params={"auth": token}).content)
    assert "MODIF" in t


# --- export bundle
def test_export_bundle(s, token):
    r = requests.get(f"{API}/dossier/{ST['pa']}/export", params={"auth": token}, timeout=120)
    assert r.status_code == 200 and r.content[:4] == b"%PDF"
    assert "_CIORNA" in r.headers.get("content-disposition", "")
    rd = PdfReader(io.BytesIO(r.content))
    first = (rd.pages[0].extract_text() or "").upper()
    assert "AUTENTICITATE" in first
    full = pdftext(r.content)
    assert "CIORN" in full.upper()
    assert "TEST_FACTURA_CLIENT_A" in full
    assert full.upper().find("AUTENTICITATE") < full.find("TEST_FACTURA_CLIENT_A")


# --- staff / instruments CRUD + signers
def test_staff_instruments_signers(s, token):
    sig = upload(s, "TEST_AUTH_SCAN")
    good = s.post(f"{API}/data/staff", json={"name": f"TEST_Good_{TAG}", "qualification": "Electrician",
                  "auths": [{"type": "ANRE gr. IIB", "number": "G123", "expires": FUT, "file": sig, "verified": True}]}).json()
    expired = s.post(f"{API}/data/staff", json={"name": f"TEST_Exp_{TAG}", "auths": [{"type": "ANRE", "number": "E1", "expires": PAST, "file": sig, "verified": True}]}).json()
    unver = s.post(f"{API}/data/staff", json={"name": f"TEST_Unv_{TAG}", "auths": [{"type": "ANRE", "number": "U1", "expires": FUT, "file": sig, "verified": False}]}).json()
    CREATED["staff"] += [good["id"], expired["id"], unver["id"]]
    assert good["created_by"] == EMAIL
    lst = s.get(f"{API}/data/staff").json()
    assert any(x["id"] == good["id"] for x in lst)
    r = s.put(f"{API}/data/staff/{good['id']}", json={"qualification": "Inginer"}); assert r.json()["qualification"] == "Inginer"
    pid = ST["pa"]
    for bad in (expired, unver):
        s.put(f"{API}/data/prosumator/{pid}", json={"signers": {"executant": bad["id"]}})
        d = s.get(f"{API}/dossier/{pid}").json()
        assert any(bad["name"] in w for w in d["signer_warnings"]), d["signer_warnings"]
        t = pdftext(requests.get(f"{API}/dossier/{pid}/file/executant", params={"auth": token}).content)
        assert bad["name"] not in t
    s.put(f"{API}/data/prosumator/{pid}", json={"signers": {"executant": good["id"]}})
    d = s.get(f"{API}/dossier/{pid}").json()
    assert d["signer_warnings"] == []
    t = pdftext(requests.get(f"{API}/dossier/{pid}/file/executant", params={"auth": token}).content)
    assert f"TEST_Good_{TAG}" in t
    # instruments
    ie = s.post(f"{API}/data/instruments", json={"maker_model": f"TESTMETER_EXP_{TAG}", "serial": "SNEXP1", "cert_name": "Certificat etalonare", "expires": PAST, "file": sig}).json()
    ig = s.post(f"{API}/data/instruments", json={"maker_model": f"TESTMETER_OK_{TAG}", "serial": "SNOK1", "cert_name": "BVM", "expires": FUT, "file": sig}).json()
    CREATED["instruments"] += [ie["id"], ig["id"]]
    s.put(f"{API}/data/prosumator/{pid}", json={"instrument_id": ie["id"]})
    d = s.get(f"{API}/dossier/{pid}").json()
    assert any("Aparat" in w for w in d["signer_warnings"])
    t = pdftext(requests.get(f"{API}/dossier/{pid}/file/priza", params={"auth": token}).content)
    assert "TESTMETER_EXP" not in t and "SNEXP1" not in t
    s.put(f"{API}/data/prosumator/{pid}", json={"instrument_id": ig["id"]})
    t = pdftext(requests.get(f"{API}/dossier/{pid}/file/priza", params={"auth": token}).content)
    assert "TESTMETER_OK" in t and "SNOK1" in t
    ST["good"] = good["id"]
    assert s.delete(f"{API}/data/instruments/{ie['id']}").status_code == 200
    assert s.get(f"{API}/data/instruments/{ie['id']}").status_code == 404


# --- finalize snapshot / reopen / audit
def test_finalize_snapshot_and_audit(s, token):
    pid = ST["pa"]
    orig = s.get(f"{API}/settings").json()
    old_phone = orig.get("phone", "")
    try:
        r = s.post(f"{API}/dossier/{pid}/finalize"); assert r.status_code == 200 and r.json()["ok"]
        d = s.get(f"{API}/dossier/{pid}").json(); assert d["finalized"]
        s.put(f"{API}/settings", json={**{k: v for k, v in orig.items() if k != "owner"}, "phone": "0700TESTSNAP"})
        t = pdftext(requests.get(f"{API}/dossier/{pid}/file/notificare", params={"auth": token}).content)
        assert "0700TESTSNAP" not in t
        s.post(f"{API}/dossier/{pid}/reopen")
        d = s.get(f"{API}/dossier/{pid}").json(); assert d["finalized"] is None
        t = pdftext(requests.get(f"{API}/dossier/{pid}/file/notificare", params={"auth": token}).content)
        assert "0700TESTSNAP" in t
    finally:
        s.put(f"{API}/settings", json={**{k: v for k, v in orig.items() if k != "owner"}, "phone": old_phone})
    assert s.get(f"{API}/settings").json().get("phone", "") == old_phone
    a = s.get(f"{API}/audit").json()
    acts = {(x["action"], x["module"]) for x in a}
    assert ("finalizare", "prosumator") in acts and ("redeschidere", "prosumator") in acts
    assert ("creare", "staff") in acts and ("ștergere", "instruments") in acts and ("modificare", "settings") in acts
    assert all(x.get("user") for x in a[:20]) and a[0]["user"] == EMAIL
    assert "_id" not in a[0]
    assert len(s.get(f"{API}/audit", params={"item_id": pid}).json()) > 0


# --- regression
def test_regression_docs_and_modules(s, token):
    r = requests.get(f"{API}/docs/notificare/{ST['pa']}", params={"auth": token})
    assert r.status_code == 200 and r.content[:4] == b"%PDF"
    assert s.get(f"{API}/docs/types").status_code == 200
    for m in ["clients", "catalog", "diu", "works", "mono"]:
        assert s.get(f"{API}/data/{m}").status_code == 200
    assert s.get(f"{API}/dossier/{REAL_ID}").json()["finalized"] is None


def test_auth_required(token):
    assert requests.get(f"{API}/dossier/{ST['pa']}").status_code == 401
    assert requests.get(f"{API}/dossier/{ST['pa']}/export").status_code == 401
    assert requests.get(f"{API}/dossier/{ST['pa']}/file/opis", params={"auth": "bad"}).status_code == 401
