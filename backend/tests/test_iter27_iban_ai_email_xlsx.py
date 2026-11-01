import os as _dns_test_os
import pytest as _dns_test_pytest
if not (_dns_test_os.environ.get("DNS_ENERGY_TEST_EMAIL") and _dns_test_os.environ.get("DNS_ENERGY_TEST_PASSWORD") and _dns_test_os.environ.get("REACT_APP_BACKEND_URL") and _dns_test_os.environ.get("DNS_ENERGY_TEST_ALLOW_MUTATION") == "1"):
    _dns_test_pytest.skip("Set staging-only DNS_ENERGY_TEST_* credentials, REACT_APP_BACKEND_URL and DNS_ENERGY_TEST_ALLOW_MUTATION=1", allow_module_level=True)

# Iteration 27: platform Settings iban/bank -> ANRE proforma transfer/PDF; admin visa AI email draft + send; visas.xlsx export
import io
import json
import uuid
from datetime import date, timedelta

import pytest
import requests
from openpyxl import load_workbook
from pypdf import PdfReader
from pymongo import MongoClient

BASE = [l.split("=", 1)[1].strip() for l in open("/app/frontend/.env") if l.startswith("REACT_APP_BACKEND_URL=")][0]
API = BASE.rstrip("/") + "/api"
env = dict(l.strip().split("=", 1) for l in open("/app/backend/.env") if "=" in l)
mdb = MongoClient(env["MONGO_URL"].strip('"'))[env["DB_NAME"].strip('"')]
RID = uuid.uuid4().hex[:6]
T = date.today()
CASES = {"past": -12, "soon": 25}
TEN = {k: f"TEST_iter27_{k}_{RID}" for k in CASES}
NOVISA = f"TEST_iter27_novisa_{RID}"
IBAN, BANK = f"RO49AAAA1B31007593840000", f"TEST Banca {RID}"
ST = {}


def login(email, pw):
    r = requests.post(f"{API}/auth/login", json={"email": email, "password": pw})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}, r.json()["access_token"]


@pytest.fixture(scope="module")
def AH():
    return login(_dns_test_os.environ["DNS_ENERGY_TEST_EMAIL"], _dns_test_os.environ["DNS_ENERGY_TEST_PASSWORD"])


@pytest.fixture(scope="module")
def H(AH):
    return AH[0]


@pytest.fixture(scope="module")
def U(H):
    email = f"sub_iter27_{RID}@example.com"
    bill = {"client_type": "pj", "company_name": "TEST I27 SRL", "cui": "RO12345678", "regcom": "J37/1/2020", "hq_address": "Str. Test 1", "county": "Vaslui",
            "locality": "Vaslui", "country": "România", "rep_name": "TEST Rep", "phone": "0700000000"}
    assert requests.post(f"{API}/auth/register", json={"email": email, "password": "Test2026!x", "name": "TEST", "billing": bill}).status_code == 200
    uid = mdb.users.find_one({"email": email})["id"]
    ST["uid"] = uid
    requests.put(f"{API}/admin/subscribers/{uid}", json={"status": "active"}, headers=H)
    return login(email, "Test2026!x")[0]


@pytest.fixture(scope="module", autouse=True)
def seed(H):
    s = requests.get(f"{API}/settings", headers=H).json()
    ST["orig"] = {"iban": s.get("iban", ""), "bank": s.get("bank", "")}
    for k, d in CASES.items():
        mdb.operators.insert_one({"tenant": TEN[k], "company": f"TEST {k} 27 SRL", "cui": f"RO27{k}{RID}", "attest_no": f"27/{k}", "contact_name": "TEST Ion Pop",
                                  "contact_email": "delivered@resend.dev", "contact_phone": "0700", "status": "aprobat", "attest_types": ["A1", "B"],
                                  "visa_date": (T + timedelta(days=d)).isoformat(), "visa_confirmed": True, "updated_at": T.isoformat()})
    mdb.operators.insert_one({"tenant": NOVISA, "company": "TEST novisa", "cui": f"RO27nv{RID}", "contact_email": "delivered@resend.dev", "status": "aprobat"})
    yield
    s = requests.get(f"{API}/settings", headers=H).json()
    s.update(ST["orig"])
    requests.put(f"{API}/settings", json=s, headers=H)
    mdb.operators.delete_many({"tenant": {"$in": list(TEN.values()) + [NOVISA]}})
    uid = ST.get("uid")
    if uid:
        for did in [d["invoice_id"] for d in mdb.anre_dossiers.find({"owner": uid})]:
            mdb.invoices.delete_many({"id": did})
        mdb.anre_dossiers.delete_many({"owner": uid})
        for c in ("subscriptions", "invoices", "billing_profiles", "billing_history", "operators"):
            mdb[c].delete_many({"tenant": uid})
        mdb.users.delete_many({"id": uid})
    for d in mdb.anre_dossiers.find({"created_by": _dns_test_os.environ["DNS_ENERGY_TEST_EMAIL"], "kind": "vizare", "history.by": "sistem", "test_iter27": True}):
        pass


# ---- Settings iban/bank -> proforma
def test_settings_save_iban_bank(H):
    s = requests.get(f"{API}/settings", headers=H).json()
    s.update({"iban": IBAN, "bank": BANK})
    r = requests.put(f"{API}/settings", json=s, headers=H)
    assert r.status_code == 200, r.text
    g = requests.get(f"{API}/settings", headers=H).json()
    assert g["iban"] == IBAN and g["bank"] == BANK


def test_admin_dossier_transfer_has_iban(H):
    r = requests.post(f"{API}/anre/dossiers", json={"kind": "vizare", "target_types": ["A1"]}, headers=H)
    assert r.status_code == 200, r.text
    d = r.json()
    ST["admin_did"], ST["admin_inv"] = d["id"], d["invoice_id"]
    assert d["transfer"]["iban"] == IBAN and d["transfer"]["bank"] == BANK
    pdf = requests.get(f"{API}/billing/invoices/{d['invoice_id']}/pdf", headers=H)
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF"), pdf.text[:200]
    txt = "".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf.content)).pages)
    assert IBAN in txt.replace(" ", "").replace("\n", "") or IBAN in txt, txt[:600]
    mdb.anre_dossiers.delete_one({"id": d["id"]})
    mdb.invoices.delete_one({"id": d["invoice_id"]})


def test_subscriber_dossier_transfer_has_platform_iban(U):
    r = requests.post(f"{API}/anre/dossiers", json={"kind": "nou", "target_types": ["A1"]}, headers=U)
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["transfer"]["iban"] == IBAN and d["transfer"]["bank"] == BANK
    pdf = requests.get(f"{API}/billing/invoices/{d['invoice_id']}/pdf", headers=U)
    txt = "".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf.content)).pages)
    assert IBAN in txt.replace("\n", ""), txt[:600]


# ---- AI email draft
def test_ai_email_streams_subject(H):
    r = requests.post(f"{API}/admin/operators/{TEN['past']}/ai-email", json={"instructions": "amintește de discuția telefonică"}, headers=H, stream=True, timeout=120)
    assert r.status_code == 200, r.text
    assert "text/event-stream" in r.headers.get("content-type", "")
    full, done = "", False
    for line in r.iter_lines(decode_unicode=True):
        if line and line.startswith("data: "):
            ev = json.loads(line[6:])
            assert "error" not in ev, ev
            full += ev.get("t", "")
            done = done or ev.get("done", False)
    assert done
    clean = full.replace("**", "").lstrip("# ").strip()
    print(full[:600])
    assert clean.lower().startswith("subiect:"), full[:200]
    assert "TEST past 27 SRL" in full or "27/past" in full
    assert "3.200" in full or "3200" in full


def test_ai_email_404_no_visa(H):
    r = requests.post(f"{API}/admin/operators/{NOVISA}/ai-email", json={}, headers=H)
    assert r.status_code == 404
    assert requests.post(f"{API}/admin/operators/TEST_nonexistent_{RID}/ai-email", json={}, headers=H).status_code == 404


def test_ai_email_403_subscriber(U):
    assert requests.post(f"{API}/admin/operators/{TEN['past']}/ai-email", json={}, headers=U).status_code == 403
    assert requests.post(f"{API}/admin/operators/{TEN['past']}/ai-email", json={}).status_code in (401, 403)


# ---- Send email
def test_send_email_validation(H):
    for b in ({"subject": "", "body": "x"}, {"subject": "x", "body": "  "}, {}):
        assert requests.post(f"{API}/admin/operators/{TEN['soon']}/email", json=b, headers=H).status_code == 400
    assert requests.post(f"{API}/admin/operators/TEST_none_{RID}/email", json={"subject": "a", "body": "b"}, headers=H).status_code == 404


def test_send_email_403_subscriber(U):
    assert requests.post(f"{API}/admin/operators/{TEN['soon']}/email", json={"subject": "a", "body": "b"}, headers=U).status_code == 403


def test_send_email_success_updates(H):
    r = requests.post(f"{API}/admin/operators/{TEN['soon']}/email", json={"subject": "TEST vizare ANRE", "body": "Bună ziua,\n\nVă scriem despre vizare.\n\nEchipa D-ENERGY"}, headers=H, timeout=60)
    assert r.status_code == 200, r.text
    assert r.json() == {"sent": True, "to": "delivered@resend.dev"}
    o = mdb.operators.find_one({"tenant": TEN["soon"]})
    assert o.get("last_contact_at")
    assert any("e-mail trimis la delivered@resend.dev: TEST vizare ANRE" in h["event"] for h in o.get("history", []))
    v = [x for x in requests.get(f"{API}/admin/operators/visas", headers=H).json() if x["tenant"] == TEN["soon"]][0]
    assert v["last_contact_at"] == o["last_contact_at"]


# ---- Excel export
def test_xlsx_matches_visas(AH):
    H, tok = AH
    r = requests.get(f"{API}/admin/operators/visas.xlsx?auth={tok}")
    assert r.status_code == 200, r.text
    assert "spreadsheetml" in r.headers["content-type"]
    assert "attachment" in r.headers.get("content-disposition", "")
    ws = load_workbook(io.BytesIO(r.content)).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][0] == "Firmă" and rows[0][-1] == "Observații apel"
    api_rows = requests.get(f"{API}/admin/operators/visas", headers=H).json()
    assert len(rows) - 1 == len(api_rows)
    assert [str(x[1]) for x in rows[1:]] == [str(v["cui"]) for v in api_rows]
    i = [x[1] for x in rows[1:]].index(f"RO27past{RID}") + 2
    assert ws.cell(i, 1).fill.fgColor.rgb.endswith("FDECEC")
    assert ws.cell(i, 5).value == -12
    j = [x[1] for x in rows[1:]].index(f"RO27soon{RID}") + 2
    assert ws.cell(j, 11).value == str(date.today())  # last_contact_at (UTC date)
    # header-auth also works
    assert requests.get(f"{API}/admin/operators/visas.xlsx", headers=H).status_code == 200


def test_xlsx_auth(U):
    assert requests.get(f"{API}/admin/operators/visas.xlsx").status_code == 401
    assert requests.get(f"{API}/admin/operators/visas.xlsx?auth=bogus").status_code == 401
    assert requests.get(f"{API}/admin/operators/visas.xlsx", headers=U).status_code == 403
