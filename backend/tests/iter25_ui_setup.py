import os as _dns_test_os
if not (_dns_test_os.environ.get("DNS_ENERGY_TEST_EMAIL") and _dns_test_os.environ.get("DNS_ENERGY_TEST_PASSWORD") and _dns_test_os.environ.get("REACT_APP_BACKEND_URL") and _dns_test_os.environ.get("DNS_ENERGY_TEST_ALLOW_MUTATION") == "1"):
    raise SystemExit("Set staging-only DNS_ENERGY_TEST_* credentials, REACT_APP_BACKEND_URL and DNS_ENERGY_TEST_ALLOW_MUTATION=1 before running this helper.")

# Setup for iteration 25 UI: subscriber with unpaid ANRE dossier + operator record (visa_date), docx template file
import io
import json
import uuid
from datetime import date, timedelta
import requests
from docx import Document
from pymongo import MongoClient

BASE = [l.split("=", 1)[1].strip() for l in open("/app/frontend/.env") if l.startswith("REACT_APP_BACKEND_URL=")][0]
API = BASE.rstrip("/") + "/api"
env = dict(l.strip().split("=", 1) for l in open("/app/backend/.env") if "=" in l)
mdb = MongoClient(env["MONGO_URL"].strip('"'))[env["DB_NAME"].strip('"')]
RID = uuid.uuid4().hex[:6]
H = {"Authorization": "Bearer " + requests.post(f"{API}/auth/login", json={"email": _dns_test_os.environ["DNS_ENERGY_TEST_EMAIL"], "password": _dns_test_os.environ["DNS_ENERGY_TEST_PASSWORD"]}).json()["access_token"]}
email = f"sub_iter25ui_{RID}@example.com"
BILL = {"client_type": "pj", "company_name": "TEST UI25 SRL", "cui": "RO12345678", "regcom": "J37/1/2020", "hq_address": "Str. Test 1", "county": "Vaslui",
        "locality": "Vaslui", "country": "România", "rep_name": "TEST Rep", "phone": "0700000000"}
assert requests.post(f"{API}/auth/register", json={"email": email, "password": "Test2026!x", "name": "TEST UI25", "billing": BILL}).status_code == 200
uid = mdb.users.find_one({"email": email})["id"]
requests.put(f"{API}/admin/subscribers/{uid}", json={"status": "active"}, headers=H)
U = {"Authorization": "Bearer " + requests.post(f"{API}/auth/login", json={"email": email, "password": "Test2026!x"}).json()["access_token"]}
cui = "RO25" + RID[:4]
op = {"company": "TEST UI25 SRL", "cui": cui, "regcom": "J1", "contact_name": "TEST C", "contact_email": "delivered@resend.dev", "contact_phone": "0700",
      "attest_no": "25/UI", "attest_types": ["A1"], "holder_name": "TEST UI25 SRL", "attest_valid": (date.today() + timedelta(days=400)).isoformat(),
      "visa_date": (date.today() + timedelta(days=45)).isoformat()}
r = requests.put(f"{API}/operator", json=op, headers=U)
assert r.status_code == 200, r.text
d = requests.post(f"{API}/anre/dossiers", json={"kind": "nou", "target_types": ["A1"]}, headers=U).json()
doc = Document()
doc.add_paragraph("Firma: {{firma}} CUI: {{cui}}")
t = doc.add_table(rows=1, cols=2)
t.rows[0].cells[0].text, t.rows[0].cells[1].text = "{{p.nr}}", "{{p.name}}"
doc.save("/tmp/iter25_tpl.docx")
print(json.dumps({"email": email, "uid": uid, "cui": cui, "did": d["id"], "inv": d["invoice_number"], "visa": op["visa_date"]}))
