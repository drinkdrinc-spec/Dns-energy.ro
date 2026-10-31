import io
import os
import re
import math
from html import escape
from datetime import datetime, timedelta
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_RIGHT
from reportlab.platypus import (BaseDocTemplate, PageTemplate, Frame, NextPageTemplate, Paragraph, Spacer, Table,
                                TableStyle, PageBreak, KeepTogether, Image)
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import cef as CEF
import schema_cef

FD = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fonts") + os.sep
pdfmetrics.registerFont(TTFont("LS", FD + "LiberationSans-Regular.ttf"))
pdfmetrics.registerFont(TTFont("LSB", FD + "LiberationSans-Bold.ttf"))
pdfmetrics.registerFontFamily("LS", normal="LS", bold="LSB", italic="LS", boldItalic="LSB")

NAVY = colors.HexColor("#0b1f3a")
GRID = colors.HexColor("#9aa7b8")
S = {
    "h": ParagraphStyle("h", fontName="LSB", fontSize=13, leading=17, alignment=TA_CENTER, spaceAfter=2),
    "sub": ParagraphStyle("sub", fontName="LS", fontSize=9.5, leading=12.5, alignment=TA_CENTER, spaceAfter=4),
    "c": ParagraphStyle("c", fontName="LSB", fontSize=10, leading=13, alignment=TA_CENTER, spaceBefore=4, spaceAfter=4),
    "p": ParagraphStyle("p", fontName="LS", fontSize=9.6, leading=13.6, alignment=TA_JUSTIFY, spaceAfter=5),
    "r": ParagraphStyle("r", fontName="LS", fontSize=9.6, leading=13, alignment=TA_RIGHT, spaceAfter=4),
    "b": ParagraphStyle("b", fontName="LSB", fontSize=9.8, leading=13.5, spaceBefore=5, spaceAfter=3),
    "li": ParagraphStyle("li", fontName="LS", fontSize=9.6, leading=13.4, leftIndent=12, bulletIndent=2, spaceAfter=2, alignment=TA_JUSTIFY),
    "small": ParagraphStyle("small", fontName="LS", fontSize=7.8, leading=10),
    "hdrl": ParagraphStyle("hdrl", fontName="LS", fontSize=8, leading=10.5),
    "hdrr": ParagraphStyle("hdrr", fontName="LS", fontSize=8, leading=10.5, alignment=TA_RIGHT),
    "cell": ParagraphStyle("cell", fontName="LS", fontSize=8.3, leading=10.5),
    "cellb": ParagraphStyle("cellb", fontName="LSB", fontSize=8.3, leading=10.5),
    "sig": ParagraphStyle("sig", fontName="LS", fontSize=9.2, leading=12.5, alignment=TA_CENTER),
}
M = "…………………"
BREAKERS = [6, 10, 16, 20, 25, 32, 40, 50, 63, 80, 100, 125, 160]
SECTION = {6: 1.5, 10: 1.5, 16: 2.5, 20: 4, 25: 4, 32: 6, 40: 10, 50: 10, 63: 16, 80: 25, 100: 35, 125: 50, 160: 70}


def num(v, default=0.0):
    try:
        return float(str(v).replace(",", ".").split()[0])
    except Exception:
        return default


def fmt(v, d=2):
    s = f"{v:.{d}f}".rstrip("0").rstrip(".")
    return s.replace(".", ",")


def pick(i):
    return next((b for b in BREAKERS if b >= i), BREAKERS[-1])


def ro_date(v):
    if not v:
        return None
    try:
        return datetime.fromisoformat(str(v)[:10]).strftime("%d.%m.%Y")
    except Exception:
        return str(v)


def sizing(r):
    tri = "mono" not in str(r.get("supply") or "Trifazat").lower()
    cur = (lambda kw: kw * 1000 / (math.sqrt(3) * 400 * 0.95)) if tri else (lambda kw: kw * 1000 / (230 * 0.95))
    inv_kw = num(r.get("inverter_kw")) or num(r.get("power"))
    br_kw = num(r.get("bransament_kw")) or num(r.get("power")) or inv_kw
    n = int(num(r.get("panel_count")))
    w = num(r.get("panel_w"))
    kwp = num(r.get("_kwp")) or n * w / 1000
    mppt = 2 if n > 12 else 1
    i_br = cur(br_kw) if br_kw else 0
    i_inv = cur(inv_kw) if inv_kw else 0
    main = pick(i_br * 1.25) if br_kw else 0
    invb = pick(i_inv * 1.25) if inv_kw else 0
    warnings = []
    if inv_kw and br_kw and inv_kw > br_kw:
        warnings.append(f"Puterea invertorului ({fmt(inv_kw)} kW) depășește puterea branșamentului ({fmt(br_kw)} kW) — este necesar spor de putere.")
    if not tri and inv_kw > 3.68:
        warnings.append("Branșament monofazat cu invertor > 3,68 kW — verificați cerințele OD (recomandat trifazat).")
    if main and invb > main:
        warnings.append("Protecția invertorului depășește protecția generală a branșamentului — verificați selectivitatea.")
    return {"tri": tri, "phases": "3P+N+PE" if tri else "1P+N+PE", "un": "3x400/230 V" if tri else "230 V",
            "br_kw": br_kw, "inv_kw": inv_kw, "i_br": i_br, "i_inv": i_inv, "main": main, "invb": invb,
            "main_sec": SECTION.get(main, 0), "inv_sec": SECTION.get(invb, 0), "poles": 4 if tri else 2, "kwp": kwp,
            "n": n, "w": w, "mppt": mppt, "per_string": math.ceil(n / mppt) if n else 0, "warnings": warnings,
            "bat_kwh": num(r.get("battery_kwh"))}


ALIASES = {"idcode": ["cnp"], "name": ["client_name"]}


CLIENT_KEYS = {"name", "idcode", "kind", "regcom", "ci_series", "rep_name", "rep_role", "hq_address", "phone", "email", "iban", "bank",
               "address", "locality", "county", "pod", "nlc", "cadastral_nr", "cf_nr", "supplier", "operator"}


class Ctx:
    """Dossier data resolved from client sheet (client fields) / dossier record -> executant settings. Missing = None."""

    def __init__(self, rec, settings, client, opis_titles=None):
        self.r, self.s, self.c = rec, settings or {}, client or {}
        self.opis_titles = opis_titles or []
        self.used = set()

    def raw(self, k):
        self.used.add(k)
        order = (self.c, self.r) if (k in CLIENT_KEYS and self.c) else (self.r, self.c)
        for src in order:
            if src.get(k) not in (None, ""):
                return src.get(k)
        for a in ALIASES.get(k, []):
            for src in order:
                if src.get(a) not in (None, ""):
                    return src.get(a)
        return None

    def v(self, k, d=M):
        x = self.raw(k)
        return escape(str(x)) if x is not None else d

    def e(self, k, d=M):
        x = self.s.get(k)
        return escape(str(x)) if x not in (None, "") else d

    @property
    def name(self): return self.v("name")
    @property
    def company(self): return self.e("company", "DNS HOSTESS SRL")
    @property
    def rep(self): return self.e("rep")
    @property
    def date(self): return ro_date(self.raw("doc_date") or self.raw("date")) or M
    @property
    def loc(self):
        parts = [self.raw("address"), self.raw("locality")]
        county = self.raw("county")
        s = ", ".join(str(p) for p in parts if p)
        if county:
            s += f", jud. {county}"
        return escape(s) if s else M
    @property
    def pj(self): return "jurid" in str(self.raw("kind") or "").lower()
    @property
    def op_short(self):
        o = str(self.raw("operator") or "")
        return escape(o.split("—")[0].split(" Sucursala")[0].strip()) if o else M
    @property
    def tot(self):
        return CEF.check(self.r.get("cef") or {}) if (self.r.get("cef") or {}).get("inverters") else None
    @property
    def cef_kw(self):
        p = self.raw("power")
        return f"{fmt(num(p))} kW" if p else M
    @property
    def power_increase(self):
        return str(self.raw("request_type") or "").strip() == "Spor de putere + prosumator"

    @property
    def request_scope(self):
        return "spor de putere și racordare a centralei electrice fotovoltaice în regim de prosumator" if self.power_increase else "racordare a centralei electrice fotovoltaice în regim de prosumator"


# ---------- blocks ----------
def B(kind, *a):
    return (kind, *a)



# ---------- PDF renderer ----------
def pdf_header(ctx):
    left = Paragraph(f"<b><font size=10>{ctx.company}</font></b><br/>{ctx.e('cui', '')} · {ctx.e('regcom', '')}<br/>Atestat ANRE {ctx.e('anre_attest', '')}", S["hdrl"])
    right = Paragraph(f"{ctx.e('phone', '')}<br/>{ctx.e('email', '')}<br/>{ctx.e('address', '')}", S["hdrr"])
    t = Table([[left, right]], colWidths=[87 * mm, 87 * mm])
    t.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, 0), 0.6, GRID), ("VALIGN", (0, 0), (-1, -1), "BOTTOM"), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
    return [t, Spacer(1, 10)]


def pdf_table(rows, widths, head, fs=None, total_mm=174):
    st_c, st_h = S["cell"], S["cellb"]
    if fs:
        st_c = ParagraphStyle("cx", parent=S["cell"], fontSize=fs, leading=fs + 2)
        st_h = ParagraphStyle("hx", parent=S["cellb"], fontSize=fs, leading=fs + 2)
    total = sum(widths)
    cw = [w / total * total_mm * mm for w in widths]
    data = [[Paragraph(str(c), st_h if (head and i == 0) or (str(c).startswith("TOTAL")) else st_c) for c in row] for i, row in enumerate(rows)]
    t = Table(data, colWidths=cw, repeatRows=1 if head else 0)
    st = [("GRID", (0, 0), (-1, -1), 0.5, GRID), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5)]
    if head:
        st.append(("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef1f5")))
    t.setStyle(TableStyle(st))
    return t


def split_sig(lines):
    i = next((k for k, x in enumerate(lines) if x.startswith(("semnătura", "numele", "(ștampil"))), len(lines))
    return lines[:i], lines[i:]


IMG_LOADER = None
_IMG_CACHE = {}


def _img(path, w, h):
    try:
        if path not in _IMG_CACHE:
            _IMG_CACHE[path] = IMG_LOADER(path)
        return Image(io.BytesIO(_IMG_CACHE[path]), width=w, height=h, kind="proportional")
    except Exception:
        return None


def sig_images(lines, s):
    txt = " ".join(lines)
    imgs = []
    if not IMG_LOADER:
        return imgs
    for nm, p in (s.get("_sigs") or {}).items():
        if nm and escape(nm) in txt:
            i = _img(p, 34 * mm, 14 * mm)
            if i:
                imgs.append(i)
            break
    company = escape(str(s.get("company") or ""))
    if s.get("_stamp_path") and "ștampil" in txt and company and company in txt:
        i = _img(s["_stamp_path"], 24 * mm, 24 * mm)
        if i:
            imgs.append(i)
    return imgs


def pdf_sign(left, right, s=None):
    def col(lines):
        if not lines:
            return ""
        top, bot = split_sig(lines)
        imgs = sig_images(lines, s or {})
        foot = f"<font size=8 color='#555555'>________________________________<br/>{'<br/>'.join(bot)}</font>"
        if not imgs:
            return Paragraph("<br/>".join(top) + "<br/><br/><br/>" + foot, S["sig"])
        row = Table([imgs], colWidths=[37 * mm] * len(imgs))
        row.setStyle(TableStyle([("ALIGN", (0, 0), (-1, -1), "CENTER"), ("VALIGN", (0, 0), (-1, -1), "BOTTOM"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
        return [Paragraph("<br/>".join(top), S["sig"]), row, Paragraph(foot, S["sig"])]
    t = Table([[col(left), col(right)]], colWidths=[87 * mm, 87 * mm])
    t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
    return KeepTogether([Spacer(1, 14), t, Spacer(1, 8)])


def to_flow(blocks, ctx):
    out = []
    for b in blocks:
        k = b[0]
        if k == "hdr":
            out += pdf_header(ctx)
        elif k == "title":
            out.append(Paragraph(b[1], S["h"]))
        elif k in ("sub", "c", "p", "r", "b", "small"):
            out.append(Paragraph(b[1], S[k]))
        elif k == "bullets":
            out += [Paragraph(x, S["li"], bulletText="•" if (len(b) < 3 or b[2]) else None) for x in b[1]]
        elif k == "table":
            out += [pdf_table(b[1], b[2], b[3], b[4] if len(b) > 4 else None), Spacer(1, 6)]
        elif k == "sign":
            out.append(pdf_sign(b[1], b[2], ctx.s))
        elif k == "pagebreak":
            out.append(PageBreak())
    return out


def _draft(c, doc):
    c.saveState()
    c.setFont("LSB", 70); c.setFillColor(colors.Color(0.8, 0.1, 0.1, alpha=0.12))
    w, h = c._pagesize
    c.translate(w / 2, h / 2); c.rotate(40); c.drawCentredString(0, 0, "CIORNĂ")
    c.restoreState()


def render(doc_keys, rec, settings, client, opis_titles=None, draft=False):
    ctx = Ctx(CEF.flatten(rec), settings, client, opis_titles)
    buf = io.BytesIO()
    pw, ph = A4
    on = _draft if draft else (lambda c, d: None)
    port = PageTemplate("port", [Frame(18 * mm, 14 * mm, pw - 36 * mm, ph - 28 * mm, id="fp")], pagesize=A4, onPage=on)
    land = PageTemplate("land", [Frame(12 * mm, 10 * mm, ph - 24 * mm, pw - 20 * mm, id="fl")], pagesize=landscape(A4), onPage=on)
    first_land = bool(doc_keys) and doc_keys[0] in ("schema_cef", "schema")
    pdf = BaseDocTemplate(buf, pageTemplates=[land, port] if first_land else [port, land], title=f"D-ENERGY — {ctx.raw('name') or ''}")
    story = []
    keys = list(doc_keys)
    if keys and not (len(keys) == 1 and keys[0] in ("schema_cef", "schema")) and "gdpr" not in keys:
        keys.append("gdpr")
    for i, k in enumerate(keys):
        if k in ("schema_cef", "schema"):
            draw, tbl = schema_cef.story(ctx, lambda t, st="p": Paragraph(t, S[st]), lambda rows, w, head=True: pdf_table(rows, [x / mm for x in w], head, None, 273), S)
            if i:
                story += [NextPageTemplate("land"), PageBreak()]
            story += draw + [PageBreak()] + tbl + [NextPageTemplate("port")]
            continue
        if i:
            story.append(PageBreak())
        blocks = DOCS[k][1](ctx)
        if k == "pv_receptie":
            seen_sign = False
            unique_blocks = []
            for block in blocks:
                if block[0] == "sign":
                    if seen_sign:
                        continue
                    seen_sign = True
                unique_blocks.append(block)
            blocks = unique_blocks
        if not any(isinstance(b, (list, tuple)) and b and b[0] == "sign" and "Instalator/Executant" in str(b) and ctx.company in str(b) and ctx.name in str(b) and ("ștampilă" in str(b).lower() or "stampila" in str(b).lower()) for b in blocks):
            blocks = blocks + [sign(ex_sig(ctx), ben_sig(ctx))]
        story += to_flow(blocks, ctx)
    pdf.build(story)
    return buf.getvalue()


# ---------- DOCX renderer ----------
def render_docx(doc_key, rec, settings, client, opis_titles=None, blocks=None, ctx=None):
    from docx import Document
    from docx.shared import Pt, Mm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.section import WD_ORIENT
    import pymupdf
    ctx = ctx or Ctx(CEF.flatten(rec), settings, client, opis_titles)
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(18)
    sec.top_margin = sec.bottom_margin = Mm(14)
    st = d.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10)

    def runs(p, text, size=None, bold=False):
        text = re.sub(r"<font[^>]*>|</font>", "", text).replace("<br/>", "\n")
        for part in re.split(r"(<b>.*?</b>)", text, flags=re.S):
            if not part:
                continue
            b = part.startswith("<b>")
            r = p.add_run(re.sub(r"</?b>", "", part).replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">").replace("&quot;", '"').replace("&#x27;", "'"))
            r.bold = b or bold
            if size:
                r.font.size = Pt(size)
        return p

    def para(text, align=None, size=None, bold=False):
        p = d.add_paragraph()
        if align:
            p.alignment = align
        return runs(p, text, size, bold)

    if isinstance(doc_key, (list, tuple)) and len(doc_key) == 1:
        doc_key = doc_key[0]
    keys = list(doc_key) if isinstance(doc_key, (list, tuple)) else None
    if blocks is None and doc_key in ("schema_cef", "schema"):
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Mm(297), Mm(210)
        pdfb = render([doc_key], rec, settings, client)
        src = pymupdf.open(stream=pdfb, filetype="pdf")
        for i, page in enumerate(src):
            if i:
                d.add_page_break()
            d.add_picture(io.BytesIO(page.get_pixmap(dpi=150).tobytes("png")), width=Mm(265))
    else:
        if blocks is None:
            blocks = []
            doc_keys = list(keys or [doc_key])
            if doc_keys and not (len(doc_keys) == 1 and doc_keys[0] in ("schema_cef", "schema")) and "gdpr" not in doc_keys:
                doc_keys.append("gdpr")
            for i, k in enumerate(doc_keys):
                if i:
                    blocks.append(B("pagebreak"))
                doc_blocks = DOCS[k][1](ctx)
                if k == "pv_receptie":
                    seen_sign = False
                    unique_blocks = []
                    for block in doc_blocks:
                        if block[0] == "sign":
                            if seen_sign:
                                continue
                            seen_sign = True
                        unique_blocks.append(block)
                    doc_blocks = unique_blocks
                if not any(isinstance(b, (list, tuple)) and b and b[0] == "sign" and "Instalator/Executant" in str(b) and ctx.company in str(b) and ctx.name in str(b) and ("ștampilă" in str(b).lower() or "stampila" in str(b).lower()) for b in doc_blocks):
                    doc_blocks = doc_blocks + [sign(ex_sig(ctx), ben_sig(ctx))]
                blocks += doc_blocks
        for b in blocks:
            k = b[0]
            if k == "hdr":
                t = d.add_table(rows=1, cols=2)
                runs(t.cell(0, 0).paragraphs[0], f"<b>{ctx.company}</b>\n{ctx.e('cui', '')} · {ctx.e('regcom', '')}\nAtestat ANRE {ctx.e('anre_attest', '')}", 8)
                pr = t.cell(0, 1).paragraphs[0]; pr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                runs(pr, f"{ctx.e('phone', '')}\n{ctx.e('email', '')}\n{ctx.e('address', '')}", 8)
                d.add_paragraph("_" * 86).runs[0].font.color.rgb = RGBColor(0x9a, 0xa7, 0xb8)
            elif k == "title":
                para(b[1], WD_ALIGN_PARAGRAPH.CENTER, 13, True)
            elif k == "sub":
                para(b[1], WD_ALIGN_PARAGRAPH.CENTER, 9.5)
            elif k == "c":
                para(b[1], WD_ALIGN_PARAGRAPH.CENTER, 10, True)
            elif k == "r":
                para(b[1], WD_ALIGN_PARAGRAPH.RIGHT)
            elif k == "b":
                para(b[1], None, 10, True)
            elif k == "small":
                para(b[1], None, 8)
            elif k == "p":
                para(b[1], WD_ALIGN_PARAGRAPH.JUSTIFY)
            elif k == "bullets":
                for x in b[1]:
                    runs(d.add_paragraph(style="List Bullet" if (len(b) < 3 or b[2]) else None), x)
            elif k == "table":
                rows = b[1]
                t = d.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"
                fs = 7 if len(rows[0]) > 8 else 8.5
                for i, row in enumerate(rows):
                    for j, c in enumerate(row):
                        runs(t.cell(i, j).paragraphs[0], str(c), fs, bold=(i == 0 and b[3]))
                d.add_paragraph()
            elif k == "sign":
                t = d.add_table(rows=1, cols=2)
                for j, lines in enumerate([b[1], b[2]]):
                    if lines:
                        p = t.cell(0, j).paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        top, bot = split_sig(lines)
                        runs(p, "\n".join(top) + "\n\n\n________________________\n" + "\n".join(bot), 9)
                d.add_paragraph()
            elif k == "pagebreak":
                d.add_page_break()
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()
