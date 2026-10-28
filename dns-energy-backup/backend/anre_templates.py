import io
import re
import uuid
from copy import deepcopy
from datetime import datetime, timezone, date

from fastapi import HTTPException, Depends, UploadFile, File, Header, Query, Response
from docx import Document
from docx.table import _Row
from pypdf import PdfReader, PdfWriter

import server as S
import anre_doc as AD

db, api, get_current_user = S.db, S.api_router, S.get_current_user
KEYS = {"cerere": "Cerere-tip ANRE", "organigrama": "Organigrama", "dotari": "Lista dotărilor (Anexa nr. 5)", "personal": "Fișa structurii de personal (Anexa nr. 7)",
        "decl_9": "Declarație pe propria răspundere (Anexa nr. 9)", "decl_10": "Declarație prelucrare date (Anexa nr. 10)"}
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PH = {"firma": "denumirea firmei", "cui": "CUI", "regcom": "nr. Registrul Comerțului", "sediu": "sediul social", "reprezentant": "reprezentant legal", "functie": "funcția reprezentantului",
      "email": "e-mail", "telefon": "telefon", "nr_atestat": "nr. atestat existent", "tipuri_atestat_existent": "tipuri atestat existent", "valabil_pana": "valabilitate atestat existent",
      "tip_solicitare": "tipul solicitării", "tipuri_solicitate": "tipuri de atestat solicitate", "data": "data generării", "lista_personal": "personal (text, câte un rând)",
      "lista_dotari": "dotări (text, câte un rând)", "p.nr / p.name / p.role / p.grade / p.license / p.license_valid / p.contract": "rând de tabel repetat pentru fiecare angajat (Word)",
      "e.nr / e.name / e.serial / e.ownership / e.calibration": "rând de tabel repetat pentru fiecare dotare (Word)"}
RX = re.compile(r"\{\{\s*([\w.]+)\s*\}\}")
now = lambda: datetime.now(timezone.utc).isoformat()


def platform_only(user):
    if user["uid"] != S.PLATFORM["id"]:
        raise HTTPException(status_code=403, detail="Doar administratorul platformei gestionează formularele oficiale")


def ctx(d):
    c = d.get("company") or {}
    st, eq = d.get("staff") or [], d.get("equipment") or []
    return {"firma": c.get("name"), "cui": c.get("cui"), "regcom": c.get("regcom"), "sediu": c.get("address"), "reprezentant": c.get("rep_name"), "functie": c.get("rep_role"),
            "email": c.get("email"), "telefon": c.get("phone"), "nr_atestat": c.get("attest_no"), "tipuri_atestat_existent": c.get("attest_types"), "valabil_pana": c.get("attest_valid"),
            "tip_solicitare": AD.TYPES.get(d.get("kind")), "tipuri_solicitate": ", ".join(d.get("target_types") or []), "data": date.today().strftime("%d.%m.%Y"),
            "lista_personal": "\n".join(f"{i + 1}. {x.get('name') or ''} — {x.get('role') or ''} — {x.get('grade') or ''} — {x.get('license') or ''} {x.get('license_valid') or ''}" for i, x in enumerate(st)),
            "lista_dotari": "\n".join(f"{i + 1}. {x.get('name') or ''} — {x.get('serial') or ''} — {x.get('ownership') or ''} — {x.get('calibration') or ''}" for i, x in enumerate(eq))}


def _sub(text, c):
    return RX.sub(lambda m: str(c[m.group(1)]) if c.get(m.group(1)) not in (None, "") else (AD.M if m.group(1) in c or m.group(1)[:2] in ("p.", "e.") else m.group(0)), text)


def _para(p, c):
    if "{{" in p.text and p.runs:
        t = _sub(p.text, c)
        p.runs[0].text = t
        for r in p.runs[1:]:
            r.text = ""


def _cells(cells, c):
    for cell in cells:
        for p in cell.paragraphs:
            _para(p, c)
        for t in cell.tables:
            _table(t, c)


def _table(t, c):
    for row in list(t.rows):
        txt = "".join(x.text for x in row.cells)
        pre = "p" if "{{p." in txt else "e" if "{{e." in txt else None
        if not pre:
            _cells(row.cells, c)
            continue
        tr = row._tr
        for i, it in enumerate(c["_" + pre] or [{}]):
            new = deepcopy(tr)
            tr.addprevious(new)
            _cells(_Row(new, t).cells, {**c, **{f"{pre}.{k}": v for k, v in it.items()}, f"{pre}.nr": i + 1})
        tr.getparent().remove(tr)


def fill_docx(data, d):
    c = {**ctx(d), "_p": d.get("staff") or [], "_e": d.get("equipment") or []}
    doc = Document(io.BytesIO(data))
    for p in doc.paragraphs:
        _para(p, c)
    for t in doc.tables:
        _table(t, c)
    for s in doc.sections:
        for part in (s.header, s.footer):
            for p in part.paragraphs:
                _para(p, c)
            for t in part.tables:
                _table(t, c)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def pdf_fields(data):
    try:
        return list((PdfReader(io.BytesIO(data)).get_fields() or {}).keys())
    except Exception:
        return []


def fill_pdf(data, d):
    r = PdfReader(io.BytesIO(data))
    fields = r.get_fields() or {}
    c = ctx(d)
    vals = {k: str(c.get(k) or "") for k in fields if k in c}
    w = PdfWriter(clone_from=r)
    for page in w.pages:
        w.update_page_form_field_values(page, vals, auto_regenerate=False)
    w.set_need_appearances_writer(True)
    out = io.BytesIO()
    w.write(out)
    return out.getvalue()


async def filled(key, d):
    t = await db.anre_templates.find_one({"key": key}, {"_id": 0})
    if not t:
        return None
    data = S.get_object(t["path"])[0]
    if t["type"] == DOCX:
        return fill_docx(data, d), DOCX, f"{key}.docx"
    if t.get("fields"):
        return fill_pdf(data, d), "application/pdf", f"{key}.pdf"
    return None


@api.get("/admin/anre/templates")
async def tpl_list(user: dict = Depends(get_current_user)):
    platform_only(user)
    have = {t["key"]: t for t in await db.anre_templates.find({}, {"_id": 0, "path": 0}).to_list(50)}
    return {"items": [{"key": k, "label": l, "template": have.get(k)} for k, l in KEYS.items()], "placeholders": PH}


@api.post("/admin/anre/templates/{key}")
async def tpl_upload(key: str, file: UploadFile = File(...), user: dict = Depends(get_current_user)):
    platform_only(user)
    if key not in KEYS:
        raise HTTPException(status_code=400, detail="Formular necunoscut")
    data, ct = await file.read(), file.content_type or ""
    if (file.filename or "").lower().endswith(".docx"):
        ct = DOCX
    ok = (ct == "application/pdf" and data.startswith(b"%PDF")) or (ct == DOCX and data.startswith(b"PK"))
    if not ok or len(data) > 15 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="Încărcați formularul oficial ca PDF sau Word (.docx), maxim 15 MB")
    if ct == DOCX:
        try:
            doc = Document(io.BytesIO(data))
            txt = "\n".join(p.text for p in doc.paragraphs) + "".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        except Exception:
            raise HTTPException(status_code=400, detail="Fișierul Word nu poate fi citit")
        fields = sorted(set(RX.findall(txt)))
    else:
        fields = pdf_fields(data)
    st = S.put_object(f"{S.APP_NAME}/anre_templates/{key}_{uuid.uuid4().hex[:6]}_{(file.filename or 'formular')[-50:]}", data, ct)
    t = {"key": key, "path": st["path"], "name": file.filename, "type": ct, "size": len(data), "fields": fields,
         "matched": [f for f in fields if f in ctx({}) or f[:2] in ("p.", "e.")], "at": now(), "by": user.get("email")}
    await db.anre_templates.update_one({"key": key}, {"$set": t}, upsert=True)
    await S.audit(user, "formular oficial ANRE încărcat", "anre", key, [file.filename or ""])
    return await tpl_list(user)


@api.delete("/admin/anre/templates/{key}")
async def tpl_delete(key: str, user: dict = Depends(get_current_user)):
    platform_only(user)
    await db.anre_templates.delete_one({"key": key})
    await S.audit(user, "formular oficial ANRE eliminat", "anre", key, [])
    return await tpl_list(user)


@api.get("/admin/anre/templates/{key}/file")
async def tpl_file(key: str, authorization: str = Header(None), auth: str = Query(None)):
    user = await S.user_from_token(authorization, auth)
    platform_only(user)
    t = await db.anre_templates.find_one({"key": key}, {"_id": 0})
    if not t:
        raise HTTPException(status_code=404, detail="Formular neîncărcat")
    return Response(content=S.get_object(t["path"])[0], media_type=t["type"], headers={"Content-Disposition": S.disp("inline", t["name"] or key), "Cache-Control": "no-store"})
